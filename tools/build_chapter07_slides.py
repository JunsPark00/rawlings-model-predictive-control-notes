"""Build the editable chapter 7 deck from reviewed content and notebook figures.

Run tools/check_notebooks.py --write for chapter 7 before refreshing the deck.
Export the resulting PPTX with tools/export_pdf.ps1, then review every PDF page.
"""

from pathlib import Path
import argparse
import json
import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageFont

parser = argparse.ArgumentParser(
    description="제7장 한글 PPT를 내용 JSON과 실습 그래프에서 재생성합니다."
)
parser.add_argument("--manifest", type=Path)
args = parser.parse_args()

ROOT = Path(__file__).resolve().parents[1]
CH = ROOT / "chapter07_explicit_control/slides"
DATA = json.loads(
    (ROOT / "tools/chapter07_slide_content.json").read_text(encoding="utf-8")
)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BG = "F5F3EC"
INK = "13293D"
TEAL = "087E82"
GRAY = "5C6B76"
ORANGE = "D96B43"
WHITE = "FFFFFF"
COLAB = "https://colab.research.google.com/github/JunsPark00/rawlings-model-predictive-control-notes/blob/main/chapter07_explicit_control/examples/chapter07.ipynb"
FONT_REG = Path(os.environ.get("MPC_SLIDE_FONT", "/mnt/c/Windows/Fonts/malgun.ttf"))
FONT_BOLD = Path(
    os.environ.get("MPC_SLIDE_FONT_BOLD", "/mnt/c/Windows/Fonts/malgunbd.ttf")
)
if not FONT_REG.exists() or not FONT_BOLD.exists():
    raise FileNotFoundError(
        "한글 글꼴 경로를 MPC_SLIDE_FONT 및 MPC_SLIDE_FONT_BOLD로 지정하세요."
    )
manifest = []


def rect(s, x, y, w, h, col):
    sh = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(col)
    sh.line.fill.background()
    return sh


def text(s, x, y, w, h, t, size=20, col=INK, bold=False, math=False):
    if not math:
        font = ImageFont.truetype(str(FONT_BOLD if bold else FONT_REG), round(size * 4))
        lines = []
        for original in t.split("\n"):
            line = ""
            for word in original.split(" "):
                trial = (line + " " + word).strip()
                if line and font.getlength(trial) > (w * 72 - 6) * 4:
                    lines.append(line)
                    line = word
                else:
                    line = trial
            lines.append(line)
        t = "\n".join(lines)
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(t.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        p.line_spacing = 1.14
        chunks = (
            re.split(r"([_^](?:\{[^}]+\}|[A-Za-z0-9]+|[α-ωΑ-Ω∞]))", line)
            if math
            else [line]
        )
        for chunk in chunks:
            if not chunk:
                continue
            r = p.add_run()
            baseline = 0
            if math and chunk[0] in "_^":
                baseline = -25000 if chunk[0] == "_" else 30000
                chunk = chunk[1:].strip("{}")
            r.text = chunk
            r.font.name = (
                "Cambria Math"
                if math and not re.search("[가-힣]", chunk)
                else "맑은 고딕"
            )
            r.font.size = Pt(size * 0.72 if baseline else size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col)
            if baseline:
                r._r.get_or_add_rPr().set("baseline", str(baseline))
    return sh


def slide(title, kind, ref="제7장", subtitle="원본 교재의 절 순서에 따른 한글 해설"):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    font = ImageFont.truetype(str(FONT_BOLD), 116)
    title_size = min(29, 29 * (12.1 * 72 - 10) * 4 / max(1, font.getlength(title)))
    text(s, 0.6, 0.36, 12, 0.3, kind, 12, TEAL, True)
    text(s, 0.6, 0.91, 12.1, 0.62, title, title_size, INK, True)
    text(s, 0.6, 1.66, 12, 0.38, subtitle, 14, GRAY)
    rect(s, 0.6, 2.19, 0.65, 0.05, ORANGE)
    text(
        s,
        0.6,
        7.12,
        11.4,
        0.2,
        f"Rawlings · Mayne · Diehl  /  2판 1쇄  /  {ref}",
        9,
        GRAY,
    )
    text(s, 12.1, 7.05, 0.6, 0.3, f"{len(prs.slides):02d}", 12, GRAY, True)
    manifest.append({"page": len(prs.slides), "title": title, "kind": kind, "ref": ref})
    return s


def columns(s, points, eqs):
    text(s, 0.7, 2.5, 5.75, 0.35, "핵심 설명", 13, TEAL, True)
    rect(s, 7.0, 2.46, 5.65, 3.84, WHITE)
    for i, p in enumerate(points):
        y = 3.02 + 1.03 * i
        text(s, 0.72, y, 0.36, 0.36, str(i + 1), 16, ORANGE, True)
        text(s, 1.16, y, 5.36, 0.87, p, 18)
    for i, (label, eq) in enumerate(eqs):
        y = 2.72 + 1.08 * i
        text(s, 7.25, y, 5.14, 0.31, label, 12, TEAL, True)
        visible = re.sub(r"[_^{}]", "", eq)
        size = 20 if len(visible) > 34 else 22
        if re.search("[가-힣]", eq):
            size = 17
        size = min(size, 22 * 30 / max(30, len(visible)))
        text(s, 7.25, y + 0.38, 5.12, 0.61, eq, size, math=True)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def derivation(title, ref, subtitle, rows):
    s = slide(title, "핵심 관계", ref, subtitle)
    for i, (label, eq, explain) in enumerate(rows):
        y = 2.6 + i * 1.22
        rect(s, 0.7, y, 11.9, 1.05, WHITE)
        text(s, 0.93, y + 0.13, 2.3, 0.65, label, 16, TEAL, True)
        text(s, 3.35, y + 0.06, 8.85, 0.48, eq, 18 if len(eq) > 55 else 23, math=True)
        text(s, 3.35, y + 0.63, 8.85, 0.3, explain, 13, GRAY)
    return s


def figure(d, item):
    s = slide(
        item["title"],
        "실습 결과",
        f"제7장 §{d['n']} · 인쇄 pp.{d['p']}",
        item["settings"],
    )
    path = CH / f"{item['fig']}.png"
    with Image.open(path) as im:
        w, h = im.size
    scale = min(11.9 / w, 3.85 / h)
    iw, ih = w * scale, h * scale
    s.shapes.add_picture(
        str(path),
        Inches((13.333 - iw) / 2),
        Inches(2.35 + (3.85 - ih) / 2),
        width=Inches(iw),
        height=Inches(ih),
    )
    text(s, 0.8, 6.35, 11.7, 0.6, item["interpret"], 15)
    notes(s, item["settings"] + "\n" + item["interpret"] + "\n\n" + d["body"])


s = slide(
    "제약 선형 시스템의 명시적 제어법칙",
    "제7장",
    "제7장 · 인쇄 pp.451–490",
    "최적화를 미리 풀고 상태 영역별 입력 식으로 제어합니다",
)
text(
    s, 0.75, 2.7, 7.5, 1.5, "어느 상태 영역에서\n어떤 입력 식을 사용할까", 33, INK, True
)
text(
    s,
    0.8,
    4.8,
    7.3,
    1.1,
    "파라미터 QP · 활성 제약 · 연속성\n구간별 이차 DP · 파라미터 LP",
    22,
    GRAY,
)
rect(s, 8.65, 2.6, 3.9, 3.5, TEAL)
text(s, 8.95, 2.95, 3.2, 0.6, "원본 교재 제7장", 20, WHITE, True)
text(
    s,
    9.0,
    3.8,
    3.1,
    1.7,
    "11개 주요 절\n18개 절·하위 절\n한글 해설과 실행 실습",
    18,
    WHITE,
)
s = slide(
    "제7장의 학습 흐름",
    "학습 지도",
    "제7장 · 절 구성",
    "원본 교재의 절 번호와 순서를 유지합니다",
)
for i, (title, body) in enumerate(
    [
        ("7.1–7.4 명시적 QP", "활성 제약에서 임계 영역과 제어법칙을 계산합니다."),
        ("7.5–7.6 동적 계획법", "구간별 이차 가치 함수를 한 단계씩 연결합니다."),
        ("7.7–7.11 LP와 구현", "유일성을 구분하고 계산·탐색·검증을 살펴봅니다."),
    ]
):
    x = 0.7 + i * 4.22
    rect(s, x, 2.68, 3.98, 3.3, WHITE)
    text(s, x + 0.2, 3.05, 3.58, 0.9, title, 20, TEAL, True)
    text(s, x + 0.2, 4.2, 3.58, 1.4, body, 18)
s = slide(
    "상태와 입력 열을 구분합니다",
    "기호와 실습 범위",
    "제7장 · 공통 안내",
    "원본 예제 7.11·7.12와 별도 보충 실습을 구분합니다",
)
columns(
    s,
    [
        "현재 상태는 파라미터이고 입력 열은 결정 변수입니다.",
        "전체 입력 열을 계산하고 첫 입력만 적용합니다.",
        "영역은 수식으로 계산하며 격자는 검증에 사용합니다.",
    ],
    [
        ("현재 상태", "x"),
        ("예측 입력 열", "U = (u_0, …, u_{N−1})"),
        ("아핀 제어법칙", "u = K_i x + k_i"),
    ],
)
for d in DATA:
    eqs = d["eqs"]
    if d["n"] == "7.3":
        eqs = [
            ("이차 항", "UᵀHU/2"),
            ("상태와 입력의 결합", "UᵀFx"),
            ("원래 제약", "GU ≤ Wx + b"),
        ]
    if d["n"] == "7.3.3":
        eqs = [
            ("정지 조건", "HU + Fx + Gᵀλ = 0"),
            ("승수 부호", "λ ≥ 0"),
            ("상보성", "λ_j s_j = 0  (s_j: 제약 여유)"),
        ]
    if d["n"] == "7.5":
        eqs = [
            ("이전 가치 함수", "V_{j−1}: 구간별 이차"),
            ("다음 상태", "y = Ax + Bu"),
            ("한 단계 최소화", "ℓ(x,u) + V_{j−1}(y)"),
        ]
    if d["n"] == "7.9":
        eqs = [
            ("제약 수 가정", "6N개"),
            ("독립 활성 행 수", "0개부터 N개까지"),
            ("현재 예측 구간", "N = 2 → 후보 79개"),
        ]
    s = slide(
        f"{d['n']}  {d['ko']}", "교재 해설", f"제7장 §{d['n']} · 인쇄 pp.{d['p']}"
    )
    columns(s, d["points"], eqs)
    notes(s, d["body"] + f"\n원본 교재 §{d['n']}, 인쇄 pp.{d['p']}.")
    if d["n"] == "7.2":
        derivation(
            "세 영역의 해를 직접 적어 봅니다",
            "제7장 §7.2 · 보충 문제",
            "J=x²+xu+u²/2,  |u|≤1 · 각 행을 영역별로 분리합니다",
            [
                (
                    "왼쪽 포화 영역",
                    "x < −1 :  u* = 1,  V* = x² + x + 1/2",
                    "제약 없는 입력 −x가 상한 1을 넘음",
                ),
                (
                    "내부 영역",
                    "−1 ≤ x ≤ 1 :  u* = −x,  V* = x²/2",
                    "입력 제약이 최적 입력을 바꾸지 않음",
                ),
                (
                    "오른쪽 포화 영역",
                    "x > 1 :  u* = −1,  V* = x² − x + 1/2",
                    "제약 없는 입력 −x가 하한 −1보다 작음",
                ),
            ],
        )
    if d["n"] == "7.3.1":
        derivation(
            "주 실습 모델과 제약을 고정합니다",
            "제7장 §7.3–7.4 · 보충 모델",
            "2상태 · 1입력 · 예측 구간 N=2 · 이차 비용에 1/2 계수 사용",
            [
                (
                    "상태 방정식",
                    "x₁⁺ = 0.8x₁ + 0.2x₂ + 0.2u",
                    "두 번째 상태: x₂⁺ = 0.7x₂ + 0.5u",
                ),
                (
                    "상태와 입력 제한",
                    "|x₁| ≤ 2.5,  |x₂| ≤ 1.5,  |u| ≤ 0.7",
                    "초기 상태 상자는 파라미터 영역에 따로 반영",
                ),
                (
                    "가중치와 종단 조건",
                    "Q = I,  R = 0.3,  AᵀPA − P = −Q",
                    "종단 상자: 각 상태 성분의 절댓값이 1 이하",
                ),
            ],
        )
    if d["n"] == "7.3.4":
        derivation(
            "큰 KKT 행렬을 두 식으로 나눕니다",
            "제7장 §7.3.4 · 인쇄 pp.462–466",
            "활성 제약을 고정하면 입력과 승수에 대한 선형 시스템입니다",
            [
                (
                    "비용의 정지 조건",
                    "HU + G_Iᵀλ_I = −Fx",
                    "입력 열과 활성 제약의 승수를 함께 계산",
                ),
                (
                    "활성 제약 등식",
                    "G_I U = W_I x + b_I",
                    "독립 활성 행만 사용해 특이한 선형 시스템을 피함",
                ),
                (
                    "파라미터별 해",
                    "U = K_i x + k_i,  λ_I = L_i x + l_i",
                    "구한 식을 원래 제약과 승수 비음수 조건에 대입",
                ),
            ],
        )
        derivation(
            "아핀 입력을 대입하면 이차 비용이 됩니다",
            "제7장 §7.3.4 · 비용 구조",
            "J=UᵀHU/2 + UᵀFx + xᵀYx/2,  U=Kx+k",
            [
                ("이차 계수", "Π = Y + KᵀHK + KᵀF + FᵀK", "상태의 이차 항을 모음"),
                (
                    "일차·상수 계수",
                    "q = KᵀHk + Fᵀk,  r = kᵀHk/2",
                    "상수 입력 k가 일차 항과 상수 항을 만듦",
                ),
                (
                    "영역별 가치 함수",
                    "V_i(x) = xᵀΠx/2 + qᵀx + r",
                    "노트북에서 원래 다단계 QP 비용과 비교",
                ),
            ],
        )
    if d["n"] == "7.4":
        derivation(
            "명시적 해에도 종단 조건이 필요합니다",
            "제7장 §7.4 · 보충 폐루프",
            "이 모델의 영 입력 종단 제어와 리아푸노프 등식에 근거합니다",
            [
                (
                    "종단 불변성",
                    "x ∈ X_f  ⇒  Ax ∈ X_f",
                    "행별 절댓값 합이 1 이하인 A와 단위 상자",
                ),
                (
                    "이동 입력 후보",
                    "U_shift = (u_1*, 0)",
                    "현재 최적 입력의 나머지 뒤에 영 입력을 추가",
                ),
                (
                    "최적 비용 감소",
                    "V(x⁺) − V(x) + ℓ(x,u) ≤ 0",
                    "다음 최적해는 실행 가능한 이동 후보보다 비싸지 않음",
                ),
            ],
        )
    if d["n"] == "7.5":
        derivation(
            "스칼라 DP 후보를 구간별로 계산합니다",
            "제7장 §7.5 · 보충 스칼라 모델",
            "x⁺=0.9x+u,  |u|≤0.4,  |x|≤2,  종단 상태 한계 0.6",
            [
                (
                    "이전 이차 조각",
                    "V(y) = Py²/2 + qy + c,  y ∈ [l,r]",
                    "다음 상태가 속할 가치 함수 조각을 하나 선택",
                ),
                (
                    "자유 최적 입력",
                    "u_free = −(0.9Px + q)/(0.2 + P)",
                    "입력·다음 상태 경계를 만족하는 구간에서만 유효",
                ),
                (
                    "경계 후보",
                    "u = ±0.4,  u = l−0.9x,  u = r−0.9x",
                    "후보의 실행 가능 구간과 미분 부호를 검사",
                ),
            ],
        )
    if d["n"] == "7.7.1":
        derivation(
            "한 솔버 출력은 유일성의 증거가 아닙니다",
            "제7장 §7.7.1 · 보충 LP",
            "min(u₁+u₂),  u₁+u₂≥t,  u₁≥0, u₂≥0",
            [
                ("첫 번째 최적해", "(u₁,u₂) = (t,0)", "목적함수 값은 t"),
                ("두 번째 최적해", "(u₁,u₂) = (0,t)", "목적함수 값은 똑같이 t"),
                (
                    "전체 최적해 집합",
                    "(u₁,u₂) = (αt,(1−α)t),  0 ≤ α ≤ 1",
                    "t>0이면 선분 위 무한히 많은 최적해가 존재",
                ),
            ],
        )
    if d["n"] == "7.8":
        derivation(
            "노름 비용을 LP로 옮깁니다",
            "제7장 §7.8 · 보충 1단계 MPC",
            "보조 변수는 절댓값 위에 놓고 양의 가중치로 최소화합니다",
            [
                ("입력 절댓값", "t_u ≥ u,  t_u ≥ −u", "최적점에서 t_u=|u|"),
                (
                    "다음 상태 절댓값",
                    "t_x ≥ 0.9x+u,  t_x ≥ −0.9x−u",
                    "최적점에서 t_x=|0.9x+u|",
                ),
                (
                    "선형 목적함수",
                    "min (|x| + 0.4t_u + 2t_x)",
                    "현재 상태 x는 고정 파라미터이므로 |x|는 상수",
                ),
            ],
        )
    if d["fig"]:
        figure(d, d)
    for item in d.get("extras", []):
        if item["kind"] == "figure":
            figure(d, item)
        else:
            s = derivation(
                item["title"],
                f"제7장 §{d['n']} · 인쇄 pp.{d['p']}",
                item["subtitle"],
                item["rows"],
            )
            notes(s, d["body"] + "\n" + item["subtitle"])
s = slide(
    "검증 결과를 읽는 순서",
    "실습 검증",
    "제7장 · 실행 결과",
    "성공한 최적화 한 번보다 여러 조건을 함께 확인합니다",
)
columns(
    s,
    [
        "원래 상태·입력 QP와 명시적 입력·비용을 비교합니다.",
        "공유 경계, 제약, 이동 후보, 비용 감소를 검사합니다.",
        "DP와 LP도 각각 독립 최적화 문제와 대조합니다.",
    ],
    [
        ("QP 검증", "실행 가능성 · KKT · 해 일치"),
        ("폐루프 검증", "제약 준수 · 시간 간 비용 감소"),
        ("DP와 LP 검증", "입력 · 비용 · 경계 · 영역 밖 판정"),
    ],
)
s = slide(
    "노트북 하나로 실행하고 확인합니다",
    "실습 시작",
    "제7장 · 실행 안내",
    "한글 설명 → 수식 → 실습 → 결과 → 검증 지표",
)
columns(
    s,
    [
        "README의 제7장 코랩 버튼을 누릅니다.",
        "공통 준비부터 위에서 아래로 모두 실행합니다.",
        "모델이나 제약을 바꾸면 검증 결과를 다시 확인합니다.",
    ],
    [
        ("노트북 구성", "48개 셀 · 코드 21개 · 그래프 18개"),
        ("계산 방식", "CPU · 고정 난수 시드"),
        ("자료 구성", "기존 폴더의 PPT · PDF · 노트북"),
    ],
)
button = rect(s, 0.85, 6.35, 3.2, 0.45, WHITE)
button.click_action.hyperlink.address = COLAB
sh = text(s, 1.1, 6.4, 2.8, 0.3, "제7장 코랩에서 실행", 14, WHITE, True)
sh.click_action.hyperlink.address = COLAB
sh.text_frame.paragraphs[0].runs[0].hyperlink.address = COLAB
prs.save(CH / "chapter07_slides.pptx")
if args.manifest:
    args.manifest.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
print("Saved", len(prs.slides), "slides")
