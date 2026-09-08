"""Build the editable chapter 2 deck from reviewed content and notebook figures.

Run tools/check_notebooks.py --write for chapter 2 before refreshing the deck.
Export the resulting PPTX with tools/export_pdf.ps1, then review every PDF page.
"""

from pathlib import Path
import argparse
import json
import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageFont

parser = argparse.ArgumentParser(
    description="제2장 한글 PPT를 내용 JSON과 실습 그래프에서 재생성합니다."
)
parser.add_argument("--manifest", type=Path)
args = parser.parse_args()

ROOT = Path(__file__).resolve().parents[1]
CH = ROOT / "chapter02_mpc_regulation/slides"
DATA = json.loads((ROOT / "tools/chapter02_slide_content.json").read_text(encoding="utf-8"))
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BG = "F5F3EC"
INK = "13293D"
TEAL = "087E82"
GRAY = "5C6B76"
ORANGE = "D96B43"
WHITE = "FFFFFF"
COLAB = "https://colab.research.google.com/github/JunsPark00/rawlings-model-predictive-control-notes/blob/main/chapter02_mpc_regulation/examples/chapter02.ipynb"
FONT_REG = Path(os.environ.get("MPC_SLIDE_FONT", "/mnt/c/Windows/Fonts/malgun.ttf"))
FONT_BOLD = Path(os.environ.get("MPC_SLIDE_FONT_BOLD", "/mnt/c/Windows/Fonts/malgunbd.ttf"))
if not FONT_REG.exists() or not FONT_BOLD.exists():
    raise FileNotFoundError(
        "한글 글꼴 경로를 MPC_SLIDE_FONT 및 MPC_SLIDE_FONT_BOLD로 지정하세요."
    )
manifest = []


def rect(s, x, y, w, h, col):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(col)
    sh.line.fill.background()
    return sh


def text(s, x, y, w, h, t, size=20, col=INK, bold=False, math=False):
    if not math:
        font = ImageFont.truetype(str(FONT_BOLD if bold else FONT_REG), round(size * 4))
        lines = []
        for original in t.split("\n"):
            line = ""
            for word in original.split(" "):
                trial = (line + " " + word).strip()
                if line and font.getlength(trial) > (w * 72 - 6) * 4:
                    lines.append(line)
                    line = word
                else:
                    line = trial
            lines.append(line)
        t = "\n".join(lines)
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(t.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        p.line_spacing = 1.14
        chunks = (
            re.split(r"([_^](?:\{[^}]+\}|[A-Za-z0-9]+|[α-ωΑ-Ω∞]))", line)
            if math
            else [line]
        )
        for chunk in chunks:
            if not chunk:
                continue
            r = p.add_run()
            baseline = 0
            if math and chunk[0] in "_^":
                baseline = -25000 if chunk[0] == "_" else 30000
                chunk = chunk[1:].strip("{}")
            r.text = chunk
            r.font.name = (
                "Cambria Math" if math and not re.search("[가-힣]", chunk) else "맑은 고딕"
            )
            r.font.size = Pt(size * 0.72 if baseline else size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col)
            if baseline:
                r._r.get_or_add_rPr().set("baseline", str(baseline))
    return sh


def slide(title, kind, ref="제2장", subtitle="원본 교재의 절 순서에 따른 한글 해설"):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    font = ImageFont.truetype(str(FONT_BOLD), 116)
    title_size = min(29, 29 * (12.1 * 72 - 10) * 4 / max(1, font.getlength(title)))
    text(s, 0.6, 0.36, 12, 0.3, kind, 12, TEAL, True)
    text(s, 0.6, 0.91, 12.1, 0.62, title, title_size, INK, True)
    text(s, 0.6, 1.66, 12, 0.38, subtitle, 14, GRAY)
    rect(s, 0.6, 2.19, 0.65, 0.05, ORANGE)
    text(
        s, 0.6, 7.12, 11.4, 0.2, f"Rawlings · Mayne · Diehl  /  2판 1쇄  /  {ref}", 9, GRAY
    )
    text(s, 12.1, 7.05, 0.6, 0.3, f"{len(prs.slides):02d}", 12, GRAY, True)
    manifest.append({"page": len(prs.slides), "title": title, "kind": kind, "ref": ref})
    return s


def columns(s, points, eqs):
    text(s, 0.7, 2.5, 5.75, 0.35, "핵심 설명", 13, TEAL, True)
    rect(s, 7.0, 2.46, 5.65, 3.84, WHITE)
    for i, p in enumerate(points):
        y = 3.02 + 1.03 * i
        text(s, 0.72, y, 0.36, 0.36, str(i + 1), 16, ORANGE, True)
        text(s, 1.16, y, 5.36, 0.87, p, 18)
    for i, (label, eq) in enumerate(eqs):
        y = 2.72 + 1.08 * i
        text(s, 7.25, y, 5.14, 0.31, label, 12, TEAL, True)
        visible = re.sub(r"[_^{}]", "", eq)
        size = 20 if len(visible) > 34 else 22
        if re.search("[가-힣]", eq):
            size = 17
        size = min(size, 22 * 30 / max(30, len(visible)))
        text(s, 7.25, y + 0.38, 5.12, 0.61, eq, size, math=True)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def derivation(title, ref, subtitle, rows):
    s = slide(title, "핵심 관계", ref, subtitle)
    for i, (label, eq, explain) in enumerate(rows):
        y = 2.6 + i * 1.22
        rect(s, 0.7, y, 11.9, 1.05, WHITE)
        text(s, 0.93, y + 0.13, 2.3, 0.65, label, 16, TEAL, True)
        text(s, 3.35, y + 0.06, 8.85, 0.48, eq, 18 if len(eq) > 55 else 23, math=True)
        text(s, 3.35, y + 0.63, 8.85, 0.3, explain, 13, GRAY)
    return s


def figure(d, item):
    s = slide(
        item["title"], "실습 결과", f"제2장 §{d['n']} · 인쇄 pp.{d['p']}", item["settings"]
    )
    path = CH / f"{item['fig']}.png"
    with Image.open(path) as im:
        w, h = im.size
    scale = min(11.9 / w, 3.85 / h)
    iw, ih = w * scale, h * scale
    s.shapes.add_picture(
        str(path),
        Inches((13.333 - iw) / 2),
        Inches(2.35 + (3.85 - ih) / 2),
        width=Inches(iw),
        height=Inches(ih),
    )
    text(s, 0.8, 6.35, 11.7, 0.6, item["interpret"], 15)
    notes(s, item["settings"] + "\n" + item["interpret"] + "\n\n" + d["body"])


s = slide(
    "모델 예측 제어 — 조절",
    "제2장 · 품질 보강판",
    "제2장 · 인쇄 pp.89–192",
    "제약을 지키며 다음 시점에도 풀 수 있고, 목표로 수렴하는 제어",
)
text(s, 0.8, 2.75, 7.2, 1.5, "실행 가능성에서\n안정성으로", 35, INK, True)
text(
    s,
    0.8,
    4.7,
    7.2,
    1.2,
    "원본 QP와 제약의 기하\n종단 집합 · 차선 제어 · 경제적 비용",
    22,
    GRAY,
)
rect(s, 8.65, 2.6, 3.9, 3.5, TEAL)
text(s, 8.95, 2.95, 3.2, 0.6, "원본 교재 제2장", 20, WHITE, True)
text(
    s,
    8.95,
    3.8,
    3.2,
    1.7,
    "27개 절·하위 절\n교재 사례와 보충 실습\n한글 설명·수식·그래프",
    18,
    WHITE,
)
s = slide(
    "원본에서 실험과 증명으로 이어지는 흐름",
    "학습 지도",
    "제2장 · 공통 안내",
    "현재 해의 최적성, 다음 문제의 실행 가능성, 폐루프 안정성을 구분합니다",
)
columns(
    s,
    [
        "2.2–2.3: 문제의 해와 실행 가능 집합을 계산합니다.",
        "2.4–2.7: 종단 설계와 이동 후보로 안정성을 연결합니다.",
        "2.8–2.9: 경제적 비용과 이산 입력의 주의점을 확인합니다.",
    ],
    [
        ("원본 재현", "Examples 2.5–2.6"),
        ("원본 구조 계산", "Example 2.8의 허용 입력"),
        ("보충 실습", "불변 집합·위상·종단 보정"),
    ],
)
s = slide(
    "수치 검증이 확인하는 범위를 먼저 정합니다",
    "표기와 검증",
    "제2장 · 공통 안내",
    "원본 예제의 비용 1/2는 유지하며, 나머지 보충 LQ 실습은 공통 계수를 생략합니다",
)
columns(
    s,
    [
        "입력 이득은 u=Kx이며 K에 음의 부호를 포함합니다.",
        "격자 관찰, 해석식 대조, 집합 검증을 구분합니다.",
        "이 장의 명목 증명은 정확한 모델과 상태를 전제로 합니다.",
    ],
    [
        ("QP 검사", "최적 상태 + 제약 잔차"),
        ("다각형 검사", "모든 꼭짓점의 한 단계 변화"),
        ("비선형 최적화", "실행 가능한 국소해 · 전역 보장 아님"),
    ],
)
for d in DATA:
    s = slide(f"{d['n']}  {d['ko']}", "교재 해설", f"제2장 §{d['n']} · 인쇄 pp.{d['p']}")
    columns(s, d["points"], d["eqs"])
    notes(s, d["body"])
    for item in d["extras"]:
        if item["kind"] == "figure":
            figure(d, item)
        else:
            s = derivation(
                item["title"],
                f"제2장 §{d['n']} · 인쇄 pp.{d['p']}",
                item["subtitle"],
                item["rows"],
            )
            notes(s, d["body"])
s = slide(
    "노트북에서 가정을 바꾸고 다시 확인하세요",
    "실행과 참고",
    "제2장 · 인쇄 pp.89–192",
    "Rawlings · Mayne · Diehl, 2nd edition, 1st printing (2017)",
)
columns(
    s,
    [
        "공통 환경 준비 코드 두 셀을 실행한 뒤 개별 실습을 실행합니다.",
        "원본 수치와 직접 계산이 다른 곳은 불일치를 명시했습니다.",
        "전체 증명과 연습문제의 해답집은 아닙니다.",
    ],
    [
        ("시작 위치", "GitHub README의 제2장 Colab"),
        ("편집과 재생성", "노트북 → 그림 → PPT → PDF"),
        ("다음 장", "외란과 불확실성에 대한 강인성"),
    ],
)
link = text(s, 0.85, 6.5, 11.5, 0.38, "Colab에서 제2장 노트북 열기", 17, TEAL, True)
link.text_frame.paragraphs[0].runs[0].hyperlink.address = COLAB
link.click_action.hyperlink.address = COLAB
prs.save(CH / "chapter02_slides.pptx")
print(f"Created {len(prs.slides)} slides: {CH}")
if args.manifest:
    args.manifest.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
