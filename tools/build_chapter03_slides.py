"""Build the editable chapter 3 deck from reviewed content and notebook figures.

Run tools/check_notebooks.py --write for chapter 3 before refreshing the deck.
Export the resulting PPTX with tools/export_pdf.ps1, then review every PDF page.
"""

from pathlib import Path
import argparse
import json
import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageFont

parser = argparse.ArgumentParser(
    description="제3장 한글 PPT를 내용 JSON과 실습 그래프에서 재생성합니다."
)
parser.add_argument("--manifest", type=Path)
args = parser.parse_args()

ROOT = Path(__file__).resolve().parents[1]
CH = ROOT / "chapter03_robust_stochastic_mpc/slides"
DATA = json.loads(
    (ROOT / "tools/chapter03_slide_content.json").read_text(encoding="utf-8")
)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BG = "F5F3EC"
INK = "13293D"
TEAL = "087E82"
GRAY = "5C6B76"
ORANGE = "D96B43"
WHITE = "FFFFFF"
COLAB = "https://colab.research.google.com/github/JunsPark00/rawlings-model-predictive-control-notes/blob/main/chapter03_robust_stochastic_mpc/examples/chapter03.ipynb"
FONT_REG = Path(os.environ.get("MPC_SLIDE_FONT", "/mnt/c/Windows/Fonts/malgun.ttf"))
FONT_BOLD = Path(
    os.environ.get("MPC_SLIDE_FONT_BOLD", "/mnt/c/Windows/Fonts/malgunbd.ttf")
)
if not FONT_REG.exists() or not FONT_BOLD.exists():
    raise FileNotFoundError(
        "한글 글꼴 경로를 MPC_SLIDE_FONT 및 MPC_SLIDE_FONT_BOLD로 지정하세요."
    )
manifest = []


def rect(s, x, y, w, h, col):
    sh = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(col)
    sh.line.fill.background()
    return sh


def text(s, x, y, w, h, t, size=20, col=INK, bold=False, math=False):
    if not math:
        font = ImageFont.truetype(str(FONT_BOLD if bold else FONT_REG), round(size * 4))
        lines = []
        for original in t.split("\n"):
            line = ""
            for word in original.split(" "):
                trial = (line + " " + word).strip()
                if line and font.getlength(trial) > (w * 72 - 6) * 4:
                    lines.append(line)
                    line = word
                else:
                    line = trial
            lines.append(line)
        t = "\n".join(lines)
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(t.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        p.line_spacing = 1.14
        chunks = (
            re.split(r"([_^](?:\{[^}]+\}|[A-Za-z0-9]+|[α-ωΑ-Ω∞]))", line)
            if math
            else [line]
        )
        for chunk in chunks:
            if not chunk:
                continue
            r = p.add_run()
            baseline = 0
            if math and chunk[0] in "_^":
                baseline = -25000 if chunk[0] == "_" else 30000
                chunk = chunk[1:].strip("{}")
            r.text = chunk
            r.font.name = (
                "Cambria Math"
                if math and not re.search("[가-힣]", chunk)
                else "맑은 고딕"
            )
            r.font.size = Pt(size * 0.72 if baseline else size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col)
            if baseline:
                r._r.get_or_add_rPr().set("baseline", str(baseline))
    return sh


def slide(title, kind, ref="제3장", subtitle="원본 교재의 절 순서에 따른 한글 해설"):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    font = ImageFont.truetype(str(FONT_BOLD), 116)
    title_size = min(29, 29 * (12.1 * 72 - 10) * 4 / max(1, font.getlength(title)))
    text(s, 0.6, 0.36, 12, 0.3, kind, 12, TEAL, True)
    text(s, 0.6, 0.91, 12.1, 0.62, title, title_size, INK, True)
    text(s, 0.6, 1.66, 12, 0.38, subtitle, 14, GRAY)
    rect(s, 0.6, 2.19, 0.65, 0.05, ORANGE)
    text(
        s,
        0.6,
        7.12,
        11.4,
        0.2,
        f"Rawlings · Mayne · Diehl  /  2판 1쇄  /  {ref}",
        9,
        GRAY,
    )
    text(s, 12.1, 7.05, 0.6, 0.3, f"{len(prs.slides):02d}", 12, GRAY, True)
    manifest.append({"page": len(prs.slides), "title": title, "kind": kind, "ref": ref})
    return s


def columns(s, points, eqs):
    text(s, 0.7, 2.5, 5.75, 0.35, "핵심 설명", 13, TEAL, True)
    rect(s, 7.0, 2.46, 5.65, 3.84, WHITE)
    for i, p in enumerate(points):
        y = 3.02 + 1.03 * i
        text(s, 0.72, y, 0.36, 0.36, str(i + 1), 16, ORANGE, True)
        text(s, 1.16, y, 5.36, 0.87, p, 18)
    for i, (label, eq) in enumerate(eqs):
        y = 2.72 + 1.08 * i
        text(s, 7.25, y, 5.14, 0.31, label, 12, TEAL, True)
        visible = re.sub(r"[_^{}]", "", eq)
        size = 20 if len(visible) > 34 else 22
        if re.search("[가-힣]", eq):
            size = 17
        size = min(size, 22 * 30 / max(30, len(visible)))
        text(s, 7.25, y + 0.38, 5.12, 0.61, eq, size, math=True)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def derivation(title, ref, subtitle, rows):
    s = slide(title, "핵심 관계", ref, subtitle)
    for i, (label, eq, explain) in enumerate(rows):
        y = 2.6 + i * 1.22
        rect(s, 0.7, y, 11.9, 1.05, WHITE)
        text(s, 0.93, y + 0.13, 2.3, 0.65, label, 16, TEAL, True)
        text(s, 3.35, y + 0.06, 8.85, 0.48, eq, 18 if len(eq) > 55 else 23, math=True)
        text(s, 3.35, y + 0.63, 8.85, 0.3, explain, 13, GRAY)
    return s


def figure(d, item):
    s = slide(
        item["title"],
        "실습 결과",
        f"제3장 §{d['n']} · 인쇄 pp.{d['p']}",
        item["settings"],
    )
    path = CH / f"{item['fig']}.png"
    with Image.open(path) as im:
        w, h = im.size
    scale = min(11.9 / w, 3.85 / h)
    iw, ih = w * scale, h * scale
    s.shapes.add_picture(
        str(path),
        Inches((13.333 - iw) / 2),
        Inches(2.35 + (3.85 - ih) / 2),
        width=Inches(iw),
        height=Inches(ih),
    )
    text(s, 0.8, 6.35, 11.7, 0.6, item["interpret"], 15)
    notes(s, item["settings"] + "\n" + item["interpret"] + "\n\n" + d["body"])


s = slide(
    "강인 및 확률적 모델 예측 제어",
    "제3장 · 품질 보강판",
    "제3장 · 인쇄 pp.193–268",
    "불확실성의 범위와 분포를 구분하고 필요한 보장을 설계합니다",
)
text(s, 0.8, 2.75, 7.2, 1.5, "불확실한 미래에서\n무엇을 보장할 것인가", 34, INK, True)
text(
    s,
    0.8,
    4.7,
    7.2,
    1.2,
    "원본 피드백 비교 · 2차원 제약 축소\n튜브 · 정보 순서 · 확률과 검증",
    22,
    GRAY,
)
rect(s, 8.65, 2.6, 3.9, 3.5, TEAL)
text(s, 8.95, 2.95, 3.2, 0.6, "원본 교재 제3장", 20, WHITE, True)
text(
    s,
    8.95,
    3.8,
    3.2,
    1.7,
    "31개 절·하위 절\n원본 재현과 보충 실습\n한글 설명·수식·그래프",
    18,
    WHITE,
)
s = slide(
    "원본에서 실습으로 이어지는 학습 흐름",
    "학습 지도",
    "제3장 · 공통 안내",
    "정책의 정보 구조, 오차 집합, 확률의 의미를 연결합니다",
)
columns(
    s,
    [
        "3.1–3.4: 불확실성과 정책의 정보 순서를 봅니다.",
        "3.5–3.6: 오차 집합과 입력 여유로 튜브를 설계합니다.",
        "3.7–3.9: 확률 제약과 표본 검증의 의미를 구분합니다.",
    ],
    [
        ("원본 재현", "피드백 비교 · Example 3.13"),
        ("해석적 검사", "불변 반경 · 외접 보정"),
        ("표본 검사", "독립 경로 · 단측 신뢰상한"),
    ],
)
s = slide(
    "실제 상태와 명목 중심을 따로 저장합니다",
    "표기와 검증",
    "제3장 · 공통 안내",
    "K에는 음의 부호를 포함합니다 · 원본 재현 비용의 1/2는 유지",
)
columns(
    s,
    [
        "실제 상태를 매번 명목 중심으로 재설정하지 않습니다.",
        "모델 오차가 외란 집합 안에 포함되는지 확인합니다.",
        "수치 실험과 모든 외란에 대한 보장을 구분합니다.",
    ],
    [
        ("상태 분해", "x = z + e"),
        ("실제 입력", "u = v + Ke"),
        ("오차 변화", "e⁺ = (A+BK)e + w"),
    ],
)
for d in DATA:
    s = slide(
        f"{d['n']}  {d['ko']}", "교재 해설", f"제3장 §{d['n']} · 인쇄 pp.{d['p']}"
    )
    columns(s, d["points"], d["eqs"])
    notes(s, d["body"])
    for item in d["extras"]:
        if item["kind"] == "figure":
            figure(d, item)
        else:
            s = derivation(
                item["title"],
                f"제3장 §{d['n']} · 인쇄 pp.{d['p']}",
                item["subtitle"],
                item["rows"],
            )
            notes(s, d["body"])
s = slide(
    "노트북에서 가정을 바꾸고 다시 확인하세요",
    "실행과 참고",
    "제3장 · 인쇄 pp.193–268",
    "Rawlings · Mayne · Diehl, 2nd edition, 1st printing (2017)",
)
columns(
    s,
    [
        "공통 환경 준비 코드 두 셀을 실행한 뒤 개별 실습을 실행합니다.",
        "원본 재현과 직접 만든 보충 실습의 범위를 표시했습니다.",
        "전체 증명과 연습문제의 해답집은 아닙니다.",
    ],
    [
        ("시작 위치", "GitHub README의 제3장 Colab"),
        ("편집과 재생성", "노트북 → 그림 → PPT → PDF"),
        ("다음 장", "상태 추정과 이동 구간 추정"),
    ],
)
link = text(s, 0.85, 6.5, 11.5, 0.38, "Colab에서 제3장 노트북 열기", 17, TEAL, True)
link.text_frame.paragraphs[0].runs[0].hyperlink.address = COLAB
link.click_action.hyperlink.address = COLAB
prs.save(CH / "chapter03_slides.pptx")
print(f"Created {len(prs.slides)} slides: {CH}")
if args.manifest:
    args.manifest.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
