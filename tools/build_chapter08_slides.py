"""Build the editable chapter 8 deck from reviewed content and notebook figures.

Run tools/check_notebooks.py --write for chapter 8 before refreshing the deck.
Export the resulting PPTX with tools/export_pdf.ps1, then review every PDF page.
"""

from pathlib import Path
import argparse
import json
import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageFont

parser = argparse.ArgumentParser(
    description="제8장 한글 PPT를 내용 JSON과 실습 그래프에서 재생성합니다."
)
parser.add_argument("--manifest", type=Path)
args = parser.parse_args()

ROOT = Path(__file__).resolve().parents[1]
CH = ROOT / "chapter08_numerical_optimal_control/slides"
DATA = json.loads(
    (ROOT / "tools/chapter08_slide_content.json").read_text(encoding="utf-8")
)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BG = "F5F3EC"
INK = "13293D"
TEAL = "087E82"
GRAY = "5C6B76"
ORANGE = "D96B43"
WHITE = "FFFFFF"
COLAB = "https://colab.research.google.com/github/JunsPark00/rawlings-model-predictive-control-notes/blob/main/chapter08_numerical_optimal_control/examples/chapter08.ipynb"
FONT_REG = Path(os.environ.get("MPC_SLIDE_FONT", "/mnt/c/Windows/Fonts/malgun.ttf"))
FONT_BOLD = Path(
    os.environ.get("MPC_SLIDE_FONT_BOLD", "/mnt/c/Windows/Fonts/malgunbd.ttf")
)
if not FONT_REG.exists() or not FONT_BOLD.exists():
    raise FileNotFoundError(
        "한글 글꼴 경로를 MPC_SLIDE_FONT 및 MPC_SLIDE_FONT_BOLD로 지정하세요."
    )
manifest = []


def rect(s, x, y, w, h, col):
    sh = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(col)
    sh.line.fill.background()
    return sh


def text(s, x, y, w, h, t, size=20, col=INK, bold=False, math=False):
    if not math:
        font = ImageFont.truetype(str(FONT_BOLD if bold else FONT_REG), round(size * 4))
        lines = []
        for original in t.split("\n"):
            line = ""
            for word in original.split(" "):
                trial = (line + " " + word).strip()
                if line and font.getlength(trial) > (w * 72 - 6) * 4:
                    lines.append(line)
                    line = word
                else:
                    line = trial
            lines.append(line)
        t = "\n".join(lines)
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(t.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        p.line_spacing = 1.14
        chunks = (
            re.split(r"([_^](?:\{[^}]+\}|[A-Za-z0-9]+|[α-ωΑ-Ω∞]))", line)
            if math
            else [line]
        )
        for chunk in chunks:
            if not chunk:
                continue
            r = p.add_run()
            baseline = 0
            if math and chunk[0] in "_^":
                baseline = -25000 if chunk[0] == "_" else 30000
                chunk = chunk[1:].strip("{}")
            r.text = chunk
            r.font.name = (
                "Cambria Math"
                if math and not re.search("[가-힣]", chunk)
                else "맑은 고딕"
            )
            r.font.size = Pt(size * 0.72 if baseline else size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col)
            if baseline:
                r._r.get_or_add_rPr().set("baseline", str(baseline))
    return sh


def slide(title, kind, ref="제8장", subtitle="원본 교재의 절 순서에 따른 한글 해설"):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    font = ImageFont.truetype(str(FONT_BOLD), 116)
    title_size = min(29, 29 * (12.1 * 72 - 10) * 4 / max(1, font.getlength(title)))
    text(s, 0.6, 0.36, 12, 0.3, kind, 12, TEAL, True)
    text(s, 0.6, 0.91, 12.1, 0.62, title, title_size, INK, True)
    text(s, 0.6, 1.66, 12, 0.38, subtitle, 14, GRAY)
    rect(s, 0.6, 2.19, 0.65, 0.05, ORANGE)
    text(
        s,
        0.6,
        7.12,
        11.4,
        0.2,
        f"Rawlings · Mayne · Diehl  /  2판 1쇄  /  {ref}",
        9,
        GRAY,
    )
    text(s, 12.1, 7.05, 0.6, 0.3, f"{len(prs.slides):02d}", 12, GRAY, True)
    manifest.append({"page": len(prs.slides), "title": title, "kind": kind, "ref": ref})
    return s


def columns(s, points, eqs):
    text(s, 0.7, 2.5, 5.75, 0.35, "핵심 설명", 13, TEAL, True)
    rect(s, 7.0, 2.46, 5.65, 3.84, WHITE)
    for i, p in enumerate(points):
        y = 3.02 + 1.03 * i
        text(s, 0.72, y, 0.36, 0.36, str(i + 1), 16, ORANGE, True)
        text(s, 1.16, y, 5.36, 0.87, p, 18)
    for i, (label, eq) in enumerate(eqs):
        y = 2.72 + 1.08 * i
        text(s, 7.25, y, 5.14, 0.31, label, 12, TEAL, True)
        visible = re.sub(r"[_^{}]", "", eq)
        size = 20 if len(visible) > 34 else 22
        if re.search("[가-힣]", eq):
            size = 17
        size = min(size, 22 * 30 / max(30, len(visible)))
        text(s, 7.25, y + 0.38, 5.12, 0.61, eq, size, math=True)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def derivation(title, ref, subtitle, rows):
    s = slide(title, "핵심 관계", ref, subtitle)
    for i, (label, eq, explain) in enumerate(rows):
        y = 2.6 + i * 1.22
        rect(s, 0.7, y, 11.9, 1.05, WHITE)
        text(s, 0.93, y + 0.13, 2.3, 0.65, label, 16, TEAL, True)
        text(s, 3.35, y + 0.06, 8.85, 0.48, eq, 18 if len(eq) > 55 else 23, math=True)
        text(s, 3.35, y + 0.63, 8.85, 0.3, explain, 13, GRAY)
    return s


def figure(d, item):
    s = slide(
        item["title"],
        "실습 결과",
        f"제8장 §{d['n']} · 인쇄 pp.{d['p']}",
        item["settings"],
    )
    path = CH / f"{item['fig']}.png"
    with Image.open(path) as im:
        w, h = im.size
    scale = min(11.9 / w, 3.85 / h)
    iw, ih = w * scale, h * scale
    s.shapes.add_picture(
        str(path),
        Inches((13.333 - iw) / 2),
        Inches(2.35 + (3.85 - ih) / 2),
        width=Inches(iw),
        height=Inches(ih),
    )
    text(s, 0.8, 6.35, 11.7, 0.6, item["interpret"], 15)
    notes(s, item["settings"] + "\n" + item["interpret"] + "\n\n" + d["body"])


import nbformat

nb = nbformat.read(CH.parent / "examples/chapter08.ipynb", 4)
cell_count = len(nb.cells)
code_count = sum(c.cell_type == "code" for c in nb.cells)
figure_count = sum(bool(d["fig"]) for d in DATA) + sum(
    e["kind"] == "figure" for d in DATA for e in d["extras"]
)
s = slide(
    "수치 최적 제어",
    "제8장",
    "제8장 · 인쇄 pp.491–600",
    "모델을 적분하고 미분하며 주어진 시간 안에서 입력을 계산합니다",
)
text(s, 0.75, 2.7, 7.5, 1.5, "최적 제어 문제를\n어떻게 실제로 풀 것인가", 33, INK, True)
text(
    s,
    0.8,
    4.8,
    7.3,
    1.1,
    "적분 · 뉴턴법 · 알고리즘 미분\n사격법 · SQP · 리카티 · 실시간 반복",
    22,
    GRAY,
)
rect(s, 8.65, 2.6, 3.9, 3.5, TEAL)
text(s, 8.95, 2.95, 3.2, 0.6, "원본 교재 제8장", 20, WHITE, True)
text(
    s,
    9.0,
    3.8,
    3.1,
    1.7,
    "11개 주요 절\n51개 절·하위 절\n한글 해설과 실행 실습",
    18,
    WHITE,
)
s = slide(
    "제8장의 학습 흐름",
    "학습 지도",
    "제8장 · 절 구성",
    "계산의 기본 도구에서 실시간 MPC로 이어집니다",
)
for i, (title, body) in enumerate(
    [
        ("8.1–8.4 계산 도구", "문제 표현, 수치 적분, 방정식, 도함수를 살펴봅니다."),
        ("8.5–8.7 최적화", "직접 이산화와 비선형 최적화 방법을 연결합니다."),
        ("8.8–8.11 온라인", "시간 구조를 활용하고 실시간 반복의 오차를 봅니다."),
    ]
):
    x = 0.7 + i * 4.22
    rect(s, x, 2.68, 3.98, 3.3, WHITE)
    text(s, x + 0.2, 3.05, 3.58, 0.9, title, 20, TEAL, True)
    text(s, x + 0.2, 4.2, 3.58, 1.4, body, 18)
s = slide(
    "모델과 오차의 종류를 구분합니다",
    "기호와 실습 범위",
    "제8장 · 공통 안내",
    "원본 예제 재현과 보충 실습의 모델·범위를 구분합니다",
)
columns(
    s,
    [
        "적분 오차와 최적화 종료 오차를 따로 확인합니다.",
        "각 절의 모델·비용·제약·초기값을 명시합니다.",
        "유한 표본 검증과 일반적인 보장을 구분합니다.",
    ],
    [
        ("연속 모델 비교", "RK4 사격법 · 중점 배치법"),
        ("선형 LQ 비교", "리카티 · 축약 · 전체 KKT"),
        ("비선형 이산 모델", "iLQR · 초기화 · 실시간 반복"),
    ],
)
for d in DATA:
    eqs = d["eqs"]
    if d["n"] == "8.3.4":
        eqs = [
            ("잔차 변환", "R_T(z) = T R(z)"),
            ("변환 야코비안", "J_T(z) = T J(z)"),
            ("정확한 전체 단계", "(TJ)Δz = −TR"),
        ]
    if d["n"] == "8.6.3":
        eqs = [
            ("최소제곱 비용", "F = rᵀr/2"),
            ("Gauss–Newton", "B = JᵀJ"),
            ("정확한 헤시안", "잔차 가중 이차 미분 추가"),
        ]
    if d["n"] == "8.8.3":
        eqs = [
            ("입력 곡률", "R + BᵀP_{k+1}B"),
            ("피드백 식", "u_k = K_k x_k"),
            ("계산 순서", "역방향 재귀 → 전진 모사"),
        ]
    s = slide(
        f"{d['n']}  {d['ko']}", "교재 해설", f"제8장 §{d['n']} · 인쇄 pp.{d['p']}"
    )
    columns(s, d["points"], eqs)
    notes(s, d["body"] + f"\n원본 교재 §{d['n']}, 인쇄 pp.{d['p']}.")
    if d["n"] == "8.1.4":
        derivation(
            "연속 시간 문제를 세 부분으로 나눕니다",
            "제8장 §8.1.4 · 인쇄 pp.498–501",
            "직접 방법은 이 세 요소를 유한 변수의 문제로 옮깁니다",
            [
                (
                    "미분 모델",
                    "ẋ(t) = f_c(x(t),u(t))",
                    "입력 구간 안에서도 상태는 계속 움직임",
                ),
                (
                    "누적 비용",
                    "J = ∫ ℓ_c(x(t),u(t)) dt + V_f(x(T))",
                    "상태의 수치 적분과 비용의 수치 구적을 함께 선택",
                ),
                ("제약 검사", "(x(t),u(t)) ∈ Z", "격자점 만족과 구간 전체 만족을 구분"),
            ],
        )
    if d["n"] == "8.2.1":
        derivation(
            "RK4의 긴 식을 계산 순서로 나눕니다",
            "제8장 §8.2.1 · 보충 설명",
            "상태 미분 f를 네 번 평가하고 마지막에 가중합을 계산합니다",
            [
                (
                    "첫 두 기울기",
                    "k₁=f(x),  k₂=f(x+hk₁/2)",
                    "초기 상태와 첫 중간 상태에서 평가",
                ),
                (
                    "다음 두 기울기",
                    "k₃=f(x+hk₂/2),  k₄=f(x+hk₃)",
                    "새 중간 상태와 구간 끝의 예상 상태에서 평가",
                ),
                (
                    "가중합 갱신",
                    "x⁺ = x + h(k₁+2k₂+2k₃+k₄)/6",
                    "높은 차수와 수치 안정성 조건은 따로 확인",
                ),
            ],
        )
    if d["n"] == "8.4.5":
        derivation(
            "전진·역방향은 같은 연쇄법칙입니다",
            "제8장 §8.4.2–8.4.5 · 보충 계산 그래프",
            "y=sin(ab)+a² · 출력은 하나, 입력은 두 개",
            [
                (
                    "전진 방향",
                    "δy = cos(ab)(bδa+aδb) + 2aδa",
                    "입력 방향을 순서대로 전달해 Jv를 계산",
                ),
                (
                    "입력 a의 기울기",
                    "∂y/∂a = b cos(ab) + 2a",
                    "곱셈과 제곱 두 경로의 기여를 합산",
                ),
                (
                    "입력 b의 기울기",
                    "∂y/∂b = a cos(ab)",
                    "스칼라 출력에서 역순 전파하면 전체 기울기를 얻음",
                ),
            ],
        )
    if d["n"] == "8.5":
        derivation(
            "사격법과 배치법에 같은 연속 문제를 줍니다",
            "제8장 §8.5 · 보충 모델",
            "T=2, 초기 상태 0.8 · 기본 구간 수 20 · 구간별 상수 입력",
            [
                (
                    "연속 동역학",
                    "ẋ = −x + 0.2x³ + u",
                    "단일·다중 사격법은 같은 RK4 상태 전이를 사용",
                ),
                (
                    "누적·종단 비용",
                    "ℓ_c = (x²+0.2u²)/2,  V_f = x²",
                    "RK4 사격법은 RK4 구적, 중점 배치는 중점 구적",
                ),
                (
                    "경로 제약",
                    "|u| ≤ 1,  |x| ≤ 1.2",
                    "격자점 외에 독립 적분의 구간 내부 표본도 검사",
                ),
            ],
        )
    if d["n"] == "8.5.3":
        s = slide(
            "세 표현을 같은 기준으로 비교합니다",
            "이산화 비교",
            "제8장 §8.5 · 보충 설명",
            "같은 이산 문제와 서로 다른 이산화를 구분합니다",
        )
        columns(
            s,
            [
                "단일·다중 사격법은 같은 RK4 문제입니다.",
                "중점 배치법은 유한 격자에서 다른 이산 문제입니다.",
                "N=10,20,40으로 세분화해 차이를 확인합니다.",
            ],
            [
                ("단일 사격법", "입력만 변수 · 상태 전진 생성"),
                ("다중 사격법", "상태도 변수 · RK4 연결 제약"),
                ("중점 배치법", "중점 동역학과 중점 비용 구적"),
            ],
        )
    if d["n"] == "8.6.2":
        derivation(
            "KKT 선형 시스템을 두 행으로 나눕니다",
            "제8장 §8.6.2 · 인쇄 pp.552–553",
            "비용 정지 조건과 제약 조건이 함께 보정됩니다",
            [
                (
                    "정지 조건 보정",
                    "BΔw + G_wᵀΔλ = −∇_w L",
                    "B는 라그랑지안 헤시안 또는 적절한 근사",
                ),
                (
                    "제약 조건 보정",
                    "G_wΔw = −G(w)",
                    "현재 점의 등식 위반을 선형화해 보정",
                ),
                (
                    "원시·쌍대 갱신",
                    "w⁺=w+αΔw,  λ⁺=λ+αΔλ",
                    "최종 원래 제약과 정지 잔차를 다시 확인",
                ),
            ],
        )
    if d["n"] == "8.6.3":
        derivation(
            "Gauss–Newton에서 생략하는 항을 봅니다",
            "제8장 §8.6.3 · 인쇄 pp.553–556",
            "최소제곱 F=rᵀr/2 · J는 잔차의 야코비안",
            [
                ("기울기", "∇F = Jᵀr", "잔차와 야코비안으로 계산"),
                (
                    "정확한 헤시안",
                    "∇²F = JᵀJ + Σ_i r_i∇²r_i",
                    "잔차가 클 때 두 번째 항이 중요한 곡률을 포함",
                ),
                (
                    "Gauss–Newton 근사",
                    "B_GN = JᵀJ",
                    "잔차 영점에서는 생략한 항이 사라짐",
                ),
            ],
        )
    if d["n"] == "8.7.2":
        derivation(
            "로그 장벽의 세 관계를 분리합니다",
            "제8장 §8.7.2 · 보충 원판 문제",
            "단위 원판 안에서 목표점과의 거리 비용을 최소화합니다",
            [
                ("양의 슬랙", "s = 1−zᵀz > 0", "모든 장벽 반복은 원판 내부에서 수행"),
                (
                    "장벽 목적함수",
                    "φ_μ(z) = F(z) − μ log(s)",
                    "μ를 줄이면서 원래 경계 최적해에 접근",
                ),
                (
                    "교란 상보성",
                    "λ = μ/s,  λs = μ",
                    "유한 μ에서 원래 상보성 λs=0과 차이가 남음",
                ),
            ],
        )
    if d["n"] == "8.8.3":
        derivation(
            "리카티 재귀의 행렬식을 분리합니다",
            "제8장 §8.8.3 · 인쇄 pp.564–566",
            "u=Kx 관례 · K에 음의 부호를 포함합니다",
            [
                (
                    "입력 방향 곡률",
                    "D_k = R + BᵀP_{k+1}B",
                    "작은 입력 차원의 선형 시스템을 구성",
                ),
                (
                    "피드백 이득",
                    "D_k K_k = −BᵀP_{k+1}A",
                    "역행렬을 만들지 않고 선형 방정식을 풂",
                ),
                (
                    "가치 곡률 갱신",
                    "P_k = Q + AᵀP_{k+1}A + AᵀP_{k+1}BK_k",
                    "종단 가치 곡률부터 시간 역순으로 계산",
                ),
            ],
        )
    if d["n"] == "8.8.6":
        derivation(
            "iLQR는 피드백을 넣어 비선형 모사를 합니다",
            "제8장 §8.8.6 · 보충 이산 모델",
            "x⁺=0.85x+0.08x³+0.3u · N=12 · 이 실습에는 강제약 없음",
            [
                (
                    "역방향 근사",
                    "동역학 일차 미분 + 비용 이차 근사",
                    "동역학 이차 미분을 생략한 Gauss–Newton 변형",
                ),
                (
                    "입력 갱신",
                    "u_new = u_old + αk + K(x_new−x_old)",
                    "피드백 항은 전진 모사에서 생긴 상태 차이에 반응",
                ),
                (
                    "수락 검사",
                    "J_new < J_old",
                    "실제 비선형 비용이 줄어드는 단계 길이를 선택",
                ),
            ],
        )
    if d["n"] == "8.9.2":
        derivation(
            "한 번의 QP와 수렴한 해를 비교합니다",
            "제8장 §8.9.2 · 보충 순차 GN 실시간 반복",
            "이산 비선형 모델 · 입력 ±1.2 · 상태 강제약 없음",
            [
                (
                    "새 측정 반영",
                    "현재 상태에서 이전 입력 열을 재모사",
                    "상태 민감도로 Gauss–Newton QP 구성",
                ),
                (
                    "한 번의 갱신",
                    "입력 제한을 포함한 QP 1회",
                    "첫 입력 적용 후 나머지 입력 열을 이동",
                ),
                (
                    "남은 최적화 오차",
                    "J_a(x) − J_b(x)",
                    "a: QP 1회, b: 수렴한 해 · 같은 현재 상태에서 비교",
                ),
            ],
        )
    if d["fig"]:
        figure(d, d)
    for item in d["extras"]:
        if item["kind"] == "figure":
            figure(d, item)
        else:
            s = derivation(
                item["title"],
                f"제8장 §{d['n']} · 인쇄 pp.{d['p']}",
                item["subtitle"],
                item["rows"],
            )
            notes(s, d["body"] + "\n" + item["subtitle"])
s = slide(
    "결과를 검증하는 순서",
    "실습 검증",
    "제8장 · 실행 결과",
    "단순한 성공 플래그보다 원래 문제의 오차를 확인합니다",
)
columns(
    s,
    [
        "적분은 정확한 해·독립 적분·격자 세분화로 봅니다.",
        "도함수와 KKT는 독립 식과 잔차로 확인합니다.",
        "폐루프 결과와 남은 최적화 오차를 함께 기록합니다.",
    ],
    [
        ("적분·미분", "차수 · 감도 · 전진/역방향 일치"),
        ("최적 제어", "해 일치 · 동역학 · 입력 제한"),
        ("온라인 반복", "비용 차이 · 입력 · 최종 상태"),
    ],
)
s = slide(
    "마지막 장의 통합 노트북을 실행합니다",
    "실습 시작",
    "제8장 · 실행 안내",
    "한글 설명 → 관계식 → 실습 → 결과 해석 → 검증 지표",
)
columns(
    s,
    [
        "README의 제8장 코랩 버튼을 누릅니다.",
        "공통 준비부터 위에서 아래로 모두 실행합니다.",
        "실습마다 바뀌는 모델과 제약 조건을 확인합니다.",
    ],
    [
        ("노트북 구성", f"{cell_count}개 셀 · 코드 {code_count}개"),
        ("결과 자료", f"그래프 {figure_count}개 · 검증 지표 75개"),
        ("실행 환경", "CPU · 최초 설치에는 인터넷 필요"),
    ],
)
button = rect(s, 0.85, 6.35, 3.2, 0.45, WHITE)
button.click_action.hyperlink.address = COLAB
sh = text(s, 1.1, 6.4, 2.8, 0.3, "제8장 코랩에서 실행", 14, TEAL, True)
sh.click_action.hyperlink.address = COLAB
text(s, 4.3, 6.42, 8, 0.3, "원본 교재: 2판 1쇄 · 제8장 인쇄 pp.491–600", 12, GRAY)
prs.save(CH / "chapter08_slides.pptx")
if args.manifest:
    args.manifest.write_text(json.dumps(manifest, ensure_ascii=False, indent=2))
print("Saved", len(prs.slides), "slides")
