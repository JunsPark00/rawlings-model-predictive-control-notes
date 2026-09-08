"""Build the editable chapter 5 deck from reviewed content and notebook figures.

Run tools/check_notebooks.py --write for chapter 5 before refreshing the deck.
Export the resulting PPTX with tools/export_pdf.ps1, then review every PDF page.
"""

from pathlib import Path
import argparse
import json
import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageFont

parser = argparse.ArgumentParser(
    description="제5장 한글 PPT를 내용 JSON과 실습 그래프에서 재생성합니다."
)
parser.add_argument("--manifest", type=Path)
args = parser.parse_args()

ROOT = Path(__file__).resolve().parents[1]
CH = ROOT / "chapter05_output_mpc/slides"
DATA = json.loads(
    (ROOT / "tools/chapter05_slide_content.json").read_text(encoding="utf-8")
)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BG = "F5F3EC"
INK = "13293D"
TEAL = "087E82"
GRAY = "5C6B76"
ORANGE = "D96B43"
WHITE = "FFFFFF"
COLAB = "https://colab.research.google.com/github/JunsPark00/rawlings-model-predictive-control-notes/blob/main/chapter05_output_mpc/examples/chapter05.ipynb"
FONT_REG = Path(os.environ.get("MPC_SLIDE_FONT", "/mnt/c/Windows/Fonts/malgun.ttf"))
FONT_BOLD = Path(
    os.environ.get("MPC_SLIDE_FONT_BOLD", "/mnt/c/Windows/Fonts/malgunbd.ttf")
)
if not FONT_REG.exists() or not FONT_BOLD.exists():
    raise FileNotFoundError(
        "한글 글꼴 경로를 MPC_SLIDE_FONT 및 MPC_SLIDE_FONT_BOLD로 지정하세요."
    )
manifest = []


def rect(s, x, y, w, h, col):
    sh = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(col)
    sh.line.fill.background()
    return sh


def text(s, x, y, w, h, t, size=20, col=INK, bold=False, math=False):
    if not math:
        font = ImageFont.truetype(str(FONT_BOLD if bold else FONT_REG), round(size * 4))
        lines = []
        for original in t.split("\n"):
            line = ""
            for word in original.split(" "):
                trial = (line + " " + word).strip()
                if line and font.getlength(trial) > (w * 72 - 6) * 4:
                    lines.append(line)
                    line = word
                else:
                    line = trial
            lines.append(line)
        t = "\n".join(lines)
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(t.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        p.line_spacing = 1.14
        chunks = (
            re.split(r"([_^](?:\{[^}]+\}|[A-Za-z0-9]+|[α-ωΑ-Ω∞]))", line)
            if math
            else [line]
        )
        for chunk in chunks:
            if not chunk:
                continue
            r = p.add_run()
            baseline = 0
            if math and chunk[0] in "_^":
                baseline = -25000 if chunk[0] == "_" else 30000
                chunk = chunk[1:].strip("{}")
            r.text = chunk
            r.font.name = (
                "Cambria Math"
                if math and not re.search("[가-힣]", chunk)
                else "맑은 고딕"
            )
            r.font.size = Pt(size * 0.72 if baseline else size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col)
            if baseline:
                r._r.get_or_add_rPr().set("baseline", str(baseline))
    return sh


def slide(title, kind, ref="제5장", subtitle="원본 교재의 절 순서에 따른 한글 해설"):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    font = ImageFont.truetype(str(FONT_BOLD), 116)
    title_size = min(29, 29 * (12.1 * 72 - 10) * 4 / max(1, font.getlength(title)))
    text(s, 0.6, 0.36, 12, 0.3, kind, 12, TEAL, True)
    text(s, 0.6, 0.91, 12.1, 0.62, title, title_size, INK, True)
    text(s, 0.6, 1.66, 12, 0.38, subtitle, 14, GRAY)
    rect(s, 0.6, 2.19, 0.65, 0.05, ORANGE)
    text(
        s,
        0.6,
        7.12,
        11.4,
        0.2,
        f"Rawlings · Mayne · Diehl  /  2판 1쇄  /  {ref}",
        9,
        GRAY,
    )
    text(s, 12.1, 7.05, 0.6, 0.3, f"{len(prs.slides):02d}", 12, GRAY, True)
    manifest.append({"page": len(prs.slides), "title": title, "kind": kind, "ref": ref})
    return s


def columns(s, points, eqs):
    text(s, 0.7, 2.5, 5.75, 0.35, "핵심 설명", 13, TEAL, True)
    rect(s, 7.0, 2.46, 5.65, 3.84, WHITE)
    for i, p in enumerate(points):
        y = 3.02 + 1.03 * i
        text(s, 0.72, y, 0.36, 0.36, str(i + 1), 16, ORANGE, True)
        text(s, 1.16, y, 5.36, 0.87, p, 18)
    for i, (label, eq) in enumerate(eqs):
        y = 2.72 + 1.08 * i
        text(s, 7.25, y, 5.14, 0.31, label, 12, TEAL, True)
        visible = re.sub(r"[_^{}]", "", eq)
        size = 20 if len(visible) > 34 else 22
        if re.search("[가-힣]", eq):
            size = 17
        size = min(size, 22 * 30 / max(30, len(visible)))
        text(s, 7.25, y + 0.38, 5.12, 0.61, eq, size, math=True)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def derivation(title, ref, subtitle, rows):
    s = slide(title, "핵심 관계", ref, subtitle)
    for i, (label, eq, explain) in enumerate(rows):
        y = 2.6 + i * 1.22
        rect(s, 0.7, y, 11.9, 1.05, WHITE)
        text(s, 0.93, y + 0.13, 2.3, 0.65, label, 16, TEAL, True)
        text(s, 3.35, y + 0.06, 8.85, 0.48, eq, 18 if len(eq) > 55 else 23, math=True)
        text(s, 3.35, y + 0.63, 8.85, 0.3, explain, 13, GRAY)
    return s


def figure(d, item):
    s = slide(
        item["title"],
        "실습 결과",
        f"제5장 §{d['n']} · 인쇄 pp.{d['p']}",
        item["settings"],
    )
    path = CH / f"{item['fig']}.png"
    with Image.open(path) as im:
        w, h = im.size
    scale = min(11.9 / w, 3.85 / h)
    iw, ih = w * scale, h * scale
    s.shapes.add_picture(
        str(path),
        Inches((13.333 - iw) / 2),
        Inches(2.35 + (3.85 - ih) / 2),
        width=Inches(iw),
        height=Inches(ih),
    )
    text(s, 0.8, 6.35, 11.7, 0.6, item["interpret"], 15)
    notes(s, item["settings"] + "\n" + item["interpret"] + "\n\n" + d["body"])


s = slide(
    "출력 피드백 MPC",
    "제5장",
    "제5장 · 인쇄 pp.339–368",
    "측정에서 출발해 제약을 지키고 목표 출력을 추적합니다",
)
text(
    s, 0.75, 2.7, 7.5, 1.5, "추정 오차까지 고려해\n어떻게 제어할 것인가", 34, INK, True
)
text(
    s,
    0.8,
    4.8,
    7.3,
    1.1,
    "두 겹 튜브 · 제약 축소 · 외란 추정\n정상상태 목표 · 실행 가능성",
    22,
    GRAY,
)
rect(s, 8.65, 2.6, 3.9, 3.5, TEAL)
text(s, 8.95, 2.95, 3.2, 0.6, "원본 교재 제5장", 20, WHITE, True)
text(
    s,
    9.0,
    3.8,
    3.1,
    1.7,
    "8개 주요 절\n16개 절·하위 절\n한글 해설과 실행 실습",
    18,
    WHITE,
)
s = slide(
    "제5장의 학습 흐름",
    "학습 지도",
    "제5장 · 절 구성",
    "교재의 절·하위 절 번호와 순서를 유지합니다",
)
for i, (title, body) in enumerate(
    [
        ("5.1–5.2 정보와 구조", "상태 추정의 불확실성을 제어 문제에 반영합니다."),
        ("5.3–5.4 출력 튜브 MPC", "두 오차 집합을 계산해 상태·입력 제약을 축소합니다."),
        (
            "5.5–5.8 목표 추적",
            "외란을 추정하고 목표를 갱신하며 성립 조건을 확인합니다.",
        ),
    ]
):
    x = 0.7 + i * 4.22
    rect(s, x, 2.68, 3.98, 3.3, WHITE)
    text(s, x + 0.2, 3.05, 3.58, 0.9, title, 20, TEAL, True)
    text(s, x + 0.2, 4.2, 3.58, 1.4, body, 18)
s = slide(
    "실제·추정·명목 상태를 구분합니다",
    "기호와 정보 시점",
    "제5장 · 인쇄 pp.339–347",
    "예측형 추정: 과거 측정으로 현재 입력을 계산하고, 현재 측정으로 다음 추정을 갱신",
)
columns(
    s,
    [
        "실제 상태는 제어기가 직접 알 수 없는 값입니다.",
        "추정 상태에는 관측기의 정보가 반영됩니다.",
        "명목 상태는 외란 없는 중심 모델로 전파합니다.",
    ],
    [
        ("실제 상태", "x = z + e + ε"),
        ("추정 상태", "x̂ = z + e"),
        ("명목 상태", "z⁺ = Az + Bv"),
    ],
)
for d in DATA:
    eqs = d["eqs"]
    if d["n"] == "5.3.2":
        eqs = [
            ("관측기", "x̂⁺ = Ax̂ + Bu + L(y−Cx̂)"),
            ("추정 오차계", "ε⁺ = (A−LC)ε + w − Lη"),
            ("보충 예의 불변 반경", "s_ε = 0.06833"),
        ]
    if d["n"] == "5.4":
        eqs = [
            ("추정 반경 갱신", "s_ε⁺ = 0.4s_ε + 0.041"),
            ("추종 반경 갱신", "s_e⁺ = 0.6s_e + 0.7(s_ε+0.03)"),
            ("상태 여유", "s_x(k) = s_ε(k) + s_e(k)"),
        ]
    if d["n"] == "5.8":
        eqs = [
            ("구간 사이 거리", "max(|a−c|, |b−d|)"),
            ("집합 이동", "d_H(z+A,A) = |z|"),
            ("안정한 구간 급수의 반경", "q / (1−λ),  0 ≤ λ < 1"),
        ]
    s = slide(
        f"{d['n']}  {d['ko']}", "교재 해설", f"제5장 §{d['n']} · 인쇄 pp.{d['p']}"
    )
    columns(s, d["points"], eqs)
    notes(s, d["body"] + f"\n원본 교재 §{d['n']}, 인쇄 pp.{d['p']}.")
    if d["n"] == "5.2":
        derivation(
            "상태와 입력의 여유는 다릅니다",
            "제5장 §5.2–5.3 · 인쇄 pp.341–348",
            "두 오차를 더한 범위와 보조 피드백의 범위를 각각 계산합니다",
            [
                ("추정 오차 집합", "ε ∈ Σ", "실제 상태와 추정 상태 사이의 차이"),
                (
                    "상태의 총 여유",
                    "Γ = Σ ⊕ S",
                    "실제 상태와 명목 상태 사이에는 두 오차가 모두 포함",
                ),
                (
                    "입력의 여유",
                    "u − v = Ke ∈ KS",
                    "입력 보정에는 추종 오차만 직접 사용",
                ),
            ],
        )
    if d["n"] == "5.3.2":
        derivation(
            "추정 오차의 반경을 한 줄씩 계산합니다",
            "제5장 §5.3.2 · 인쇄 pp.344–346",
            "보충 스칼라 예 · |a−Lc|=0.4 · |w|≤0.02 · |η|≤0.03",
            [
                (
                    "오차 외란의 범위",
                    "q = 0.02 + 0.7 × 0.03 = 0.041",
                    "과정 외란과 관측기를 거친 측정 잡음을 합산",
                ),
                (
                    "불변 반경",
                    "s = q / (1−0.4) = 0.06833…",
                    "현재 오차가 이 구간 안이면 다음 오차도 구간 안",
                ),
                ("한 단계 확인", "0.4s + q = s", "구간 꼭짓점 전수 검사와 같은 최댓값"),
            ],
        )
    if d["n"] == "5.3.4":
        derivation(
            "명목 중심을 계속 전파합니다",
            "제5장 §5.3.4 · 인쇄 pp.348–351",
            "교재 알고리즘 5.5의 구조 · 추정 상태로 중심을 매번 재설정하지 않음",
            [
                (
                    "최적화와 입력",
                    "v = κ_N(z),  u = v + K(x̂−z)",
                    "명목 상태에서 MPC를 풀고 실제 입력에 오차 피드백을 추가",
                ),
                (
                    "관측기 갱신",
                    "x̂⁺ = Ax̂ + Bu + L(y−Cx̂)",
                    "실제 적용 입력과 현재 측정으로 다음 추정을 계산",
                ),
                (
                    "명목 중심 갱신",
                    "z⁺ = Az + Bv",
                    "명목 입력을 사용해 외란 없는 중심을 전파",
                ),
            ],
        )
        s = slide(
            "원점 수렴과 집합 접근을 구분합니다",
            "안정성 해석",
            "제5장 §5.3.4 · 인쇄 pp.350–352",
            "명목 MPC의 안정성과 유계 오차 집합을 함께 사용합니다",
        )
        columns(
            s,
            [
                "종단 비용·종단 집합·초기 실행 가능성을 확인합니다.",
                "명목 상태의 값 함수 감소를 검사합니다.",
                "잡음이 지속되면 실제 상태에는 오차 범위가 남습니다.",
            ],
            [
                ("명목 비용 감소", "V(z⁺) − V(z) ≤ −ℓ(z,v)"),
                ("명목 상태", "z → 0"),
                ("실제 상태", "dist(x,Γ) → 0"),
            ],
        )
    if d["n"] == "5.5.2":
        derivation(
            "외란을 추정한 뒤 목표 입력을 보정합니다",
            "제5장 §5.5.2 · 인쇄 pp.358–359",
            "보충 예 · a=0.8 · b=0.5 · 출력 설정값 1 · 외란 추정 0.12",
            [
                ("목표 상태", "x_s = 1", "보충 모델은 출력이 상태와 같음"),
                (
                    "평형 등식",
                    "0.2 × 1 = 0.5u_s + 0.12",
                    "목표 상태를 유지하는 데 필요한 입력 계산",
                ),
                (
                    "목표 입력",
                    "u_s = 0.16",
                    "외란을 무시한 목표 입력 0.4와 차이가 발생",
                ),
            ],
        )
    if d["n"] == "5.5.3":
        derivation(
            "목표를 달성할 수 있는지 먼저 확인합니다",
            "제5장 §5.5.2–5.5.3 · 인쇄 pp.358–363",
            "보충 스칼라 목표 · 목표 상태 1 · 입력 한계 0.8",
            [
                (
                    "외란 추정 0.12",
                    "u_s = (0.2−0.12)/0.5 = 0.16",
                    "평형 입력이 허용 범위 안에 있음",
                ),
                (
                    "외란 추정 0.8",
                    "u_s = (0.2−0.8)/0.5 = −1.2",
                    "입력 한계를 넘어 같은 설정값의 평형이 실행 불가능",
                ),
                (
                    "해석",
                    "목표 평형의 존재 ≠ 재귀적 실행 가능성",
                    "목표를 달성할 수 있어도 미래 MPC의 해 존재는 별도 조건",
                ),
            ],
        )
    if d["fig"]:
        figure(d, d)
    for item in d.get("extras", []):
        if item["kind"] == "figure":
            figure(d, item)
        else:
            s = derivation(
                item["title"],
                f"제5장 §{d['n']} · 인쇄 pp.{d['p']}",
                item["subtitle"],
                item["rows"],
            )
            notes(s, d["body"] + "\n" + item["subtitle"])
s = slide(
    "조건을 바꾸고 보장이 유지되는지 봅니다",
    "보충 과제",
    "제5장 §5.8 · 인쇄 p.366",
    "노트북에 여섯 가지 확장 과제와 힌트를 제공합니다",
)
columns(
    s,
    [
        "관측기 이득과 잡음 경계를 함께 바꾸어 봅니다.",
        "중심 재설정과 초기 오차 가정을 구분합니다.",
        "외란 추정·목표 계산·제약의 관계를 확인합니다.",
    ],
    [
        ("잡음과 추정", "L, w_max, η_max"),
        ("초기 집합", "ε_0 ∈ Σ,  e_0 ∈ S"),
        ("목표 입력", "|u_s| ≤ u_max"),
    ],
)
s = slide(
    "노트북 하나로 실행하고 확인합니다",
    "실습 시작",
    "제5장 · 실행 안내",
    "한글 설명 → 수식 → 보충 실습 → 결과 → 검증 지표",
)
columns(
    s,
    [
        "README에서 제5장 코랩 버튼을 누릅니다.",
        "공통 준비부터 위에서 아래로 모두 실행합니다.",
        "잡음·초기 오차·목표 조건을 바꿔 결과를 비교합니다.",
    ],
    [
        ("노트북 구성", "44개 셀 · 코드 19개 · 그래프 15개"),
        ("핵심 검증", "불변 구간 · 제약 · 오프셋 · 목표"),
        ("실행 환경", "CPU · 최초 설치와 글꼴 다운로드 필요"),
    ],
)
button = rect(s, 0.85, 6.35, 3.2, 0.45, WHITE)
button.click_action.hyperlink.address = COLAB
sh = text(s, 1.1, 6.4, 2.8, 0.3, "제5장 코랩에서 실행", 14, WHITE, True)
sh.click_action.hyperlink.address = COLAB
sh.text_frame.paragraphs[0].runs[0].hyperlink.address = COLAB
sh.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(WHITE)
prs.save(CH / "chapter05_slides.pptx")
if args.manifest:
    args.manifest.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
print("Saved", len(prs.slides), "slides")
