"""Build the editable chapter 6 deck from reviewed content and notebook figures.

Run tools/check_notebooks.py --write for chapter 6 before refreshing the deck.
Export the resulting PPTX with tools/export_pdf.ps1, then review every PDF page.
"""

from pathlib import Path
import argparse
import json
import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageFont

parser = argparse.ArgumentParser(
    description="제6장 한글 PPT를 내용 JSON과 실습 그래프에서 재생성합니다."
)
parser.add_argument("--manifest", type=Path)
args = parser.parse_args()

ROOT = Path(__file__).resolve().parents[1]
CH = ROOT / "chapter06_distributed_mpc/slides"
DATA = json.loads(
    (ROOT / "tools/chapter06_slide_content.json").read_text(encoding="utf-8")
)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BG = "F5F3EC"
INK = "13293D"
TEAL = "087E82"
GRAY = "5C6B76"
ORANGE = "D96B43"
WHITE = "FFFFFF"
COLAB = "https://colab.research.google.com/github/JunsPark00/rawlings-model-predictive-control-notes/blob/main/chapter06_distributed_mpc/examples/chapter06.ipynb"
FONT_REG = Path(os.environ.get("MPC_SLIDE_FONT", "/mnt/c/Windows/Fonts/malgun.ttf"))
FONT_BOLD = Path(
    os.environ.get("MPC_SLIDE_FONT_BOLD", "/mnt/c/Windows/Fonts/malgunbd.ttf")
)
if not FONT_REG.exists() or not FONT_BOLD.exists():
    raise FileNotFoundError(
        "한글 글꼴 경로를 MPC_SLIDE_FONT 및 MPC_SLIDE_FONT_BOLD로 지정하세요."
    )
manifest = []


def rect(s, x, y, w, h, col):
    sh = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(col)
    sh.line.fill.background()
    return sh


def text(s, x, y, w, h, t, size=20, col=INK, bold=False, math=False):
    if not math:
        font = ImageFont.truetype(str(FONT_BOLD if bold else FONT_REG), round(size * 4))
        lines = []
        for original in t.split("\n"):
            line = ""
            for word in original.split(" "):
                trial = (line + " " + word).strip()
                if line and font.getlength(trial) > (w * 72 - 6) * 4:
                    lines.append(line)
                    line = word
                else:
                    line = trial
            lines.append(line)
        t = "\n".join(lines)
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(t.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        p.line_spacing = 1.14
        chunks = (
            re.split(r"([_^](?:\{[^}]+\}|[A-Za-z0-9]+|[α-ωΑ-Ω∞]))", line)
            if math
            else [line]
        )
        for chunk in chunks:
            if not chunk:
                continue
            r = p.add_run()
            baseline = 0
            if math and chunk[0] in "_^":
                baseline = -25000 if chunk[0] == "_" else 30000
                chunk = chunk[1:].strip("{}")
            r.text = chunk
            r.font.name = (
                "Cambria Math"
                if math and not re.search("[가-힣]", chunk)
                else "맑은 고딕"
            )
            r.font.size = Pt(size * 0.72 if baseline else size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col)
            if baseline:
                r._r.get_or_add_rPr().set("baseline", str(baseline))
    return sh


def slide(title, kind, ref="제6장", subtitle="원본 교재의 절 순서에 따른 한글 해설"):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    font = ImageFont.truetype(str(FONT_BOLD), 116)
    title_size = min(29, 29 * (12.1 * 72 - 10) * 4 / max(1, font.getlength(title)))
    text(s, 0.6, 0.36, 12, 0.3, kind, 12, TEAL, True)
    text(s, 0.6, 0.91, 12.1, 0.62, title, title_size, INK, True)
    text(s, 0.6, 1.66, 12, 0.38, subtitle, 14, GRAY)
    rect(s, 0.6, 2.19, 0.65, 0.05, ORANGE)
    text(
        s,
        0.6,
        7.12,
        11.4,
        0.2,
        f"Rawlings · Mayne · Diehl  /  2판 1쇄  /  {ref}",
        9,
        GRAY,
    )
    text(s, 12.1, 7.05, 0.6, 0.3, f"{len(prs.slides):02d}", 12, GRAY, True)
    manifest.append({"page": len(prs.slides), "title": title, "kind": kind, "ref": ref})
    return s


def columns(s, points, eqs):
    text(s, 0.7, 2.5, 5.75, 0.35, "핵심 설명", 13, TEAL, True)
    rect(s, 7.0, 2.46, 5.65, 3.84, WHITE)
    for i, p in enumerate(points):
        y = 3.02 + 1.03 * i
        text(s, 0.72, y, 0.36, 0.36, str(i + 1), 16, ORANGE, True)
        text(s, 1.16, y, 5.36, 0.87, p, 18)
    for i, (label, eq) in enumerate(eqs):
        y = 2.72 + 1.08 * i
        text(s, 7.25, y, 5.14, 0.31, label, 12, TEAL, True)
        visible = re.sub(r"[_^{}]", "", eq)
        size = 20 if len(visible) > 34 else 22
        if re.search("[가-힣]", eq):
            size = 17
        size = min(size, 22 * 30 / max(30, len(visible)))
        text(s, 7.25, y + 0.38, 5.12, 0.61, eq, size, math=True)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def derivation(title, ref, subtitle, rows):
    s = slide(title, "핵심 관계", ref, subtitle)
    for i, (label, eq, explain) in enumerate(rows):
        y = 2.6 + i * 1.22
        rect(s, 0.7, y, 11.9, 1.05, WHITE)
        text(s, 0.93, y + 0.13, 2.3, 0.65, label, 16, TEAL, True)
        text(s, 3.35, y + 0.06, 8.85, 0.48, eq, 18 if len(eq) > 55 else 23, math=True)
        text(s, 3.35, y + 0.63, 8.85, 0.3, explain, 13, GRAY)
    return s


def figure(d, item):
    s = slide(
        item["title"],
        "실습 결과",
        f"제6장 §{d['n']} · 인쇄 pp.{d['p']}",
        item["settings"],
    )
    path = CH / f"{item['fig']}.png"
    with Image.open(path) as im:
        w, h = im.size
    scale = min(11.9 / w, 3.85 / h)
    iw, ih = w * scale, h * scale
    s.shapes.add_picture(
        str(path),
        Inches((13.333 - iw) / 2),
        Inches(2.35 + (3.85 - ih) / 2),
        width=Inches(iw),
        height=Inches(ih),
    )
    text(s, 0.8, 6.35, 11.7, 0.6, item["interpret"], 15)
    notes(s, item["settings"] + "\n" + item["interpret"] + "\n\n" + d["body"])


s = slide(
    "분산 MPC",
    "제6장",
    "제6장 · 인쇄 pp.369–450",
    "여러 제어기가 나누어 계산하고 전체 시스템을 제어합니다",
)
text(
    s,
    0.75,
    2.7,
    7.5,
    1.5,
    "각자 계산한 입력을\n어떻게 함께 사용할 것인가",
    34,
    INK,
    True,
)
text(
    s,
    0.8,
    4.8,
    7.3,
    1.1,
    "분권 · 비협조 · 협조 · 중앙집중\n입력 계획 공유 · 유한 반복 · 비선형 비용",
    22,
    GRAY,
)
rect(s, 8.65, 2.6, 3.9, 3.5, TEAL)
text(s, 8.95, 2.95, 3.2, 0.6, "원본 교재 제6장", 20, WHITE, True)
text(
    s,
    9.0,
    3.8,
    3.1,
    1.7,
    "7개 주요 절\n23개 절·하위 절\n한글 해설과 실행 실습",
    18,
    WHITE,
)
s = slide(
    "제6장의 학습 흐름",
    "학습 지도",
    "제6장 · 절 구성",
    "교재의 절·하위 절 번호와 순서를 유지합니다",
)
for i, (title, body) in enumerate(
    [
        ("6.1–6.2 계산과 목적", "선형 MPC의 해법과 네 가지 제어 방식을 비교합니다."),
        (
            "6.3–6.4 제약과 참여자",
            "독립·결합 제약, 추정·외란, 다수 참여자로 확장합니다.",
        ),
        (
            "6.5–6.7 비선형과 안정성",
            "후보의 실제 비용을 검사하고 시간 간 감소를 연결합니다.",
        ),
    ]
):
    x = 0.7 + i * 4.22
    rect(s, x, 2.68, 3.98, 3.3, WHITE)
    text(s, x + 0.2, 3.05, 3.58, 0.9, title, 20, TEAL, True)
    text(s, x + 0.2, 4.2, 3.58, 1.4, body, 18)
s = slide(
    "시간 단계와 최적화 반복을 구분합니다",
    "기호와 계산 순서",
    "제6장 · 인쇄 pp.375–380",
    "한 시간 안에서 여러 참여자가 입력 계획을 개선한 뒤 첫 입력을 적용합니다",
)
columns(
    s,
    [
        "실제 시스템은 시간 단계마다 한 번 진행합니다.",
        "각 시간 안에서 최적화 반복을 여러 번 할 수 있습니다.",
        "입력 계획은 다음 시간의 시작점으로 이어집니다.",
    ],
    [
        ("실제 시간", "k = 0, 1, 2, …"),
        ("최적화 반복", "p = 0, 1, …, p_max"),
        ("예측 구간", "U = (u(0), …, u(N−1))"),
    ],
)
for d in DATA:
    eqs = d["eqs"]
    if d["n"] == "6.2.4":
        eqs = [
            ("공통 목적", "J = ρ_1 J_1 + ρ_2 J_2"),
            ("두 전체 후보", "c_1=(u_1*,u_2), c_2=(u_1,u_2*)"),
            ("가중 결합", "u⁺ = w_1 c_1 + w_2 c_2"),
        ]
    if d["n"] == "6.2.6":
        eqs = [
            ("국소 관측기의 정보", "국소 측정 + 두 실제 입력"),
            ("알려진 입력 상쇄", "e_i⁺ = (A_i−L_i C_i)e_i"),
            ("정보 누락", "상대 입력 누락 → 모델 오차"),
        ]
    if d["n"] == "6.5.3":
        eqs = [
            ("보충 동역학", "f(x,u) = f_0(x) + u"),
            ("입력 제한", "|u_i| ≤ 0.6"),
            ("종단 비용", "V_f(x) = c|x|², c ≈ 2.5536"),
        ]
    if d["n"] == "6.5.4":
        eqs = [
            ("이동 시작점의 감소", "V_w⁺ ≤ V − ℓ(x,u)"),
            ("분산 반복의 개선", "V_new⁺ ≤ V_w⁺"),
            ("시간 간 비용 감소", "V_new⁺ − V ≤ −ℓ(x,u)"),
        ]
    s = slide(
        f"{d['n']}  {d['ko']}", "교재 해설", f"제6장 §{d['n']} · 인쇄 pp.{d['p']}"
    )
    columns(s, d["points"], eqs)
    notes(s, d["body"] + f"\n원본 교재 §{d['n']}, 인쇄 pp.{d['p']}.")
    if d["n"] == "6.1.1":
        derivation(
            "선형 MPC를 입력 열의 이차 문제로 만듭니다",
            "제6장 §6.1.1 · 인쇄 pp.370–375",
            "상태를 제거해도 원래 동역학과 비용은 같습니다",
            [
                (
                    "상태 열의 표현",
                    "X = S_x x_0 + S_u U",
                    "모든 예측 상태를 초기 상태와 입력 열로 표현",
                ),
                (
                    "입력 비용",
                    "J(U) = UᵀHU/2 + (Fx_0)ᵀU + const.",
                    "상태 비용을 대입해 입력 열에 대한 이차 함수 구성",
                ),
                (
                    "최적 입력 열",
                    "U* = −H⁻¹Fx_0",
                    "실제 코드는 역행렬을 만들지 않고 선형 방정식을 풂",
                ),
            ],
        )
    if d["n"] == "6.1.2":
        derivation(
            "이전 입력 열도 제어기의 기억입니다",
            "제6장 §6.1.2 · 인쇄 pp.375–380",
            "같은 상태에서도 남아 있는 입력 계획이 다르면 다음 입력이 달라질 수 있습니다",
            [
                (
                    "현재 입력 계획",
                    "U = (u(0), u(1), …, u(N−1))",
                    "최적화를 끝내지 않았다면 이전 계획의 영향이 남음",
                ),
                (
                    "다음 시작점",
                    "U_w⁺ = (u(1), …, u(N−1), 0)",
                    "이 실습은 영 입력 종단 비용 조건을 사용",
                ),
                (
                    "분석의 상태",
                    "(x,U)",
                    "현재 상태의 수렴과 균일 리아푸노프 안정성을 구분",
                ),
            ],
        )
    if d["n"] == "6.2.3":
        s = slide(
            "세 가지 질문은 서로 다릅니다",
            "수렴과 안정성",
            "제6장 §6.2.2–6.2.3 · 인쇄 pp.383–392",
            "같은 수치 예에서도 각각 다른 행렬이나 조건을 검사합니다",
        )
        columns(
            s,
            [
                "각자의 최적성 조건을 동시에 만족하는 점이 있는지 봅니다.",
                "계획을 반복 교환했을 때 그 점에 도달하는지 봅니다.",
                "계산한 입력을 실제 시스템에 적용한 결과를 봅니다.",
            ],
            [
                ("내시점의 존재", "연립 최적성 조건의 해"),
                ("반복의 수렴", "ρ(T_BR) < 1"),
                ("폐루프 안정성", "ρ(A+BK) < 1"),
            ],
        )
    if d["n"] == "6.2.4":
        derivation(
            "새 입력을 붙이는 대신 전체 후보를 평균합니다",
            "제6장 §6.2.4 · 인쇄 pp.392–398",
            "현재 입력 (0,0) · 각 참여자의 단독 개선 입력이 1인 보충 예",
            [
                (
                    "참여자 1의 전체 후보",
                    "c_1 = (1,0)",
                    "참여자 2의 입력은 이전 값으로 유지",
                ),
                (
                    "참여자 2의 전체 후보",
                    "c_2 = (0,1)",
                    "참여자 1의 입력은 이전 값으로 유지",
                ),
                (
                    "절반씩 가중한 후보",
                    "(c_1+c_2)/2 = (0.5,0.5)",
                    "두 새 입력을 그대로 붙인 (1,1)과 다른 갱신",
                ),
            ],
        )
    if d["n"] == "6.3.2":
        derivation(
            "공유 자원은 함께 재배분해야 할 수 있습니다",
            "제6장 §6.3.2 · 인쇄 pp.411–413",
            "보충 비용 J=(u₁−1)²+(u₂−2)² · u₁,u₂≥0 · u₁+u₂≤1",
            [
                (
                    "멈춘 경계점",
                    "u = (0.8,0.2),  J = 3.28",
                    "다른 입력을 고정하면 각자 더 개선할 수 없음",
                ),
                (
                    "함께 이동한 최적점",
                    "u* = (0,1),  J* = 2",
                    "참여자 1의 자원을 참여자 2에게 옮기면 전체 비용 감소",
                ),
                (
                    "핵심 구분",
                    "블록별 개선 정지 ≠ 전체 최적성",
                    "실행 가능성과 비용 감소가 있어도 최적성은 별도",
                ),
            ],
        )
    if d["n"] == "6.5.1":
        derivation(
            "비볼록 비용은 세 항으로 나누어 봅니다",
            "제6장 §6.5.1 · 인쇄 pp.421–423",
            "교재 그림 6.7의 함수에 상수 2를 더함 · 최적점과 비용 차이는 동일",
            [
                (
                    "개별 입력 항",
                    "q(u) = exp(−2u) − 2exp(−u)",
                    "각 입력의 매끄러운 비선형 비용",
                ),
                (
                    "결합 봉우리",
                    "b(u₁,u₂) = 1.1exp(−0.4[(u₁+0.2)²+(u₂+0.2)²])",
                    "두 입력이 함께 결정하는 추가 항",
                ),
                (
                    "전체 비용",
                    "J = 2 + q(u₁) + q(u₂) + b(u₁,u₂)",
                    "개별 후보 사이에 더 높은 비용 영역이 있을 수 있음",
                ),
            ],
        )
    if d["n"] == "6.5.2":
        derivation(
            "평균 후보를 검사하고 필요하면 후보를 제외합니다",
            "제6장 §6.5.2 · 인쇄 pp.423–425",
            "각 블록의 투영 기울기와 Armijo 탐색 후 전체 후보를 공유합니다",
            [
                ("후보 평균", "u_test = Σ_i w_i c_i", "남은 후보들의 가중치 합은 1"),
                (
                    "비용 검사",
                    "J(u_test) ≤ Σ_i w_i J(c_i)",
                    "비볼록 목적함수의 실제 값을 다시 계산",
                ),
                (
                    "실패한 경우",
                    "비용 개선이 가장 작은 후보를 제외",
                    "가중치를 다시 정하고 검사 · 필요하면 후보 하나만 선택",
                ),
            ],
        )
    if d["n"] == "6.5.3":
        derivation(
            "보충 비선형 모델의 종단 조건을 확인합니다",
            "제6장 §6.5.3–6.5.4 · 인쇄 pp.425–430",
            "영 입력이 가능하고 상태 강제약이 없는 보충 시스템",
            [
                (
                    "영 입력 동역학",
                    "f₀(x) = 0.7x + 0.08tanh(swap(x))",
                    "두 상태를 교환한 벡터에 tanh를 성분별로 적용",
                ),
                (
                    "전역 수축 경계",
                    "|f₀(x)| ≤ 0.78|x|",
                    "삼각 부등식과 |tanh(x)|≤|x|를 사용",
                ),
                (
                    "종단 비용 계수",
                    "c = 1/(1−0.78²)",
                    "c|f₀(x)|²−c|x|² ≤ −|x|²이므로 영 입력을 끝에 붙임",
                ),
            ],
        )
    if d["fig"]:
        figure(d, d)
    for item in d.get("extras", []):
        if item["kind"] == "figure":
            figure(d, item)
        else:
            s = derivation(
                item["title"],
                f"제6장 §{d['n']} · 인쇄 pp.{d['p']}",
                item["subtitle"],
                item["rows"],
            )
            notes(s, d["body"] + "\n" + item["subtitle"])
s = slide(
    "반복 횟수와 가정을 함께 바꾸어 봅니다",
    "보충 과제",
    "제6장 §6.7 · 인쇄 pp.435–447",
    "노트북에 여섯 가지 확장 과제와 힌트를 제공합니다",
)
columns(
    s,
    [
        "결합 크기에 따라 전체 폐루프 고유값을 비교합니다.",
        "가중 평균과 공유 자원 재배분의 차이를 확인합니다.",
        "비볼록 후보 검사와 종단 조건을 바꿀 때의 결과를 봅니다.",
    ],
    [
        ("모델 결합", "B_12, B_21"),
        ("협조 갱신", "Σ_i w_i = 1"),
        ("시간 간 감소", "이동 시작점 + 분산 개선"),
    ],
)
s = slide(
    "노트북 하나로 실행하고 확인합니다",
    "실습 시작",
    "제6장 · 실행 안내",
    "한글 설명 → 수식 → 실습 → 결과 → 검증 지표",
)
columns(
    s,
    [
        "README에서 제6장 코랩 버튼을 누릅니다.",
        "공통 준비부터 위에서 아래로 모두 실행합니다.",
        "시간 단계와 최적화 반복을 구분해 결과를 읽습니다.",
    ],
    [
        ("노트북 구성", "55개 셀 · 코드 23개 · 그래프 21개"),
        ("검증 범위", "해 일치 · 수렴 · 제약 · 비용 감소"),
        ("실행 방식", "CPU · 참여자 계산을 한 노트북에서 모사"),
    ],
)
button = rect(s, 0.85, 6.35, 3.2, 0.45, WHITE)
button.click_action.hyperlink.address = COLAB
sh = text(s, 1.1, 6.4, 2.8, 0.3, "제6장 코랩에서 실행", 14, WHITE, True)
sh.click_action.hyperlink.address = COLAB
sh.text_frame.paragraphs[0].runs[0].hyperlink.address = COLAB
prs.save(CH / "chapter06_slides.pptx")
if args.manifest:
    args.manifest.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
print("Saved", len(prs.slides), "slides")
