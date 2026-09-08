"""Build the editable chapter 1 deck from reviewed content and notebook figures.

Run tools/check_notebooks.py --write for chapter 1 before refreshing the deck.
Export the resulting PPTX with tools/export_pdf.ps1, then review every PDF page.
"""

from pathlib import Path
import argparse
import json
import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageFont

parser = argparse.ArgumentParser(
    description="제1장 한글 PPT를 내용 JSON과 실습 그래프에서 재생성합니다."
)
parser.add_argument("--manifest", type=Path)
args = parser.parse_args()

ROOT = Path(__file__).resolve().parents[1]
CH = ROOT / "chapter01_getting_started/slides"
DATA = json.loads((ROOT / "tools/chapter01_slide_content.json").read_text(encoding="utf-8"))
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BG = "F5F3EC"
INK = "13293D"
TEAL = "087E82"
GRAY = "5C6B76"
ORANGE = "D96B43"
WHITE = "FFFFFF"
COLAB = "https://colab.research.google.com/github/JunsPark00/rawlings-model-predictive-control-notes/blob/main/chapter01_getting_started/examples/chapter01.ipynb"
FONT_REG = Path(os.environ.get("MPC_SLIDE_FONT", "/mnt/c/Windows/Fonts/malgun.ttf"))
FONT_BOLD = Path(os.environ.get("MPC_SLIDE_FONT_BOLD", "/mnt/c/Windows/Fonts/malgunbd.ttf"))
if not FONT_REG.exists() or not FONT_BOLD.exists():
    raise FileNotFoundError(
        "한글 글꼴 경로를 MPC_SLIDE_FONT 및 MPC_SLIDE_FONT_BOLD로 지정하세요."
    )
manifest = []


def rect(s, x, y, w, h, col):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(col)
    sh.line.fill.background()
    return sh


def text(s, x, y, w, h, t, size=20, col=INK, bold=False, math=False):
    if not math:
        font = ImageFont.truetype(str(FONT_BOLD if bold else FONT_REG), round(size * 4))
        lines = []
        for original in t.split("\n"):
            line = ""
            for word in original.split(" "):
                trial = (line + " " + word).strip()
                if line and font.getlength(trial) > (w * 72 - 6) * 4:
                    lines.append(line)
                    line = word
                else:
                    line = trial
            lines.append(line)
        t = "\n".join(lines)
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(t.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        p.line_spacing = 1.14
        chunks = (
            re.split(r"([_^](?:\{[^}]+\}|[A-Za-z0-9]+|[α-ωΑ-Ω∞]))", line)
            if math
            else [line]
        )
        for chunk in chunks:
            if not chunk:
                continue
            r = p.add_run()
            baseline = 0
            if math and chunk[0] in "_^":
                baseline = -25000 if chunk[0] == "_" else 30000
                chunk = chunk[1:].strip("{}")
            r.text = chunk
            r.font.name = (
                "Cambria Math" if math and not re.search("[가-힣]", chunk) else "맑은 고딕"
            )
            r.font.size = Pt(size * 0.72 if baseline else size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col)
            if baseline:
                r._r.get_or_add_rPr().set("baseline", str(baseline))
    return sh


def slide(title, kind, ref="제1장", subtitle="원본 교재의 절 순서에 따른 한글 해설"):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    font = ImageFont.truetype(str(FONT_BOLD), 116)
    title_size = min(29, 29 * (12.1 * 72 - 10) * 4 / max(1, font.getlength(title)))
    text(s, 0.6, 0.36, 12, 0.3, kind, 12, TEAL, True)
    text(s, 0.6, 0.91, 12.1, 0.62, title, title_size, INK, True)
    text(s, 0.6, 1.66, 12, 0.38, subtitle, 14, GRAY)
    rect(s, 0.6, 2.19, 0.65, 0.05, ORANGE)
    text(
        s, 0.6, 7.12, 11.4, 0.2, f"Rawlings · Mayne · Diehl  /  2판 1쇄  /  {ref}", 9, GRAY
    )
    text(s, 12.1, 7.05, 0.6, 0.3, f"{len(prs.slides):02d}", 12, GRAY, True)
    manifest.append({"page": len(prs.slides), "title": title, "kind": kind, "ref": ref})
    return s


def columns(s, points, eqs):
    text(s, 0.7, 2.5, 5.75, 0.35, "핵심 설명", 13, TEAL, True)
    rect(s, 7.0, 2.46, 5.65, 3.84, WHITE)
    for i, p in enumerate(points):
        y = 3.02 + 1.03 * i
        text(s, 0.72, y, 0.36, 0.36, str(i + 1), 16, ORANGE, True)
        text(s, 1.16, y, 5.36, 0.87, p, 18)
    for i, (label, eq) in enumerate(eqs):
        y = 2.72 + 1.08 * i
        text(s, 7.25, y, 5.14, 0.31, label, 12, TEAL, True)
        visible = re.sub(r"[_^{}]", "", eq)
        size = 20 if len(visible) > 34 else 22
        if re.search("[가-힣]", eq):
            size = 17
        size = min(size, 22 * 30 / max(30, len(visible)))
        text(s, 7.25, y + 0.38, 5.12, 0.61, eq, size, math=True)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def derivation(title, ref, subtitle, rows):
    s = slide(title, "핵심 관계", ref, subtitle)
    for i, (label, eq, explain) in enumerate(rows):
        y = 2.6 + i * 1.22
        rect(s, 0.7, y, 11.9, 1.05, WHITE)
        text(s, 0.93, y + 0.13, 2.3, 0.65, label, 16, TEAL, True)
        text(s, 3.35, y + 0.06, 8.85, 0.48, eq, 18 if len(eq) > 55 else 23, math=True)
        text(s, 3.35, y + 0.63, 8.85, 0.3, explain, 13, GRAY)
    return s


def figure(d, item):
    s = slide(
        item["title"], "실습 결과", f"제1장 §{d['n']} · 인쇄 pp.{d['p']}", item["settings"]
    )
    path = CH / f"{item['fig']}.png"
    with Image.open(path) as im:
        w, h = im.size
    scale = min(11.9 / w, 3.85 / h)
    iw, ih = w * scale, h * scale
    s.shapes.add_picture(
        str(path),
        Inches((13.333 - iw) / 2),
        Inches(2.35 + (3.85 - ih) / 2),
        width=Inches(iw),
        height=Inches(ih),
    )
    text(s, 0.8, 6.35, 11.7, 0.6, item["interpret"], 15)
    notes(s, item["settings"] + "\n" + item["interpret"] + "\n\n" + d["body"])


s = slide(
    "모델 예측 제어 시작하기",
    "제1장 · 품질 보강판",
    "제1장 · 인쇄 pp.1–88",
    "원본의 구조를 따라 읽고, 코드와 그림으로 핵심 연결을 확인합니다",
)
text(
    s, 0.8, 2.7, 7.2, 1.6, "모델에서 추정과 제어까지\n직접 계산하며 이해하기", 32, INK, True
)
text(
    s,
    0.8,
    4.8,
    7.2,
    1.1,
    "초기조건 · 제약 · 동적 계획법\n칼만 필터 · MHE · 무오프셋 추종",
    22,
    GRAY,
)
rect(s, 8.65, 2.6, 3.9, 3.5, TEAL)
text(s, 8.95, 2.95, 3.2, 0.6, "원본 교재 제1장", 20, WHITE, True)
text(
    s,
    8.95,
    3.8,
    3.2,
    1.7,
    "26개 절·하위 절\n한글 설명과 실행 코드\n해석해·잔차·계수 검증",
    18,
    WHITE,
)
s = slide(
    "이 장을 읽는 순서",
    "학습 지도",
    "제1장 · 공통 안내",
    "원본 사례와 직접 설계한 보충 실습을 구분합니다",
)
columns(
    s,
    [
        "1.2: 모델의 가정과 제약을 정의합니다.",
        "1.3–1.4: 미래 입력과 현재 상태를 계산합니다.",
        "1.5: 목표와 외란을 포함해 전체 구조를 연결합니다.",
    ],
    [
        ("그림의 역할", "식을 실제 응답과 연결"),
        ("코드의 역할", "동일성·잔차·가정을 검증"),
        ("자료의 범위", "전 증명·연습문제 해답집 아님"),
    ],
)
s = slide(
    "추정·목표·제어를 하나의 순환으로 연결합니다",
    "전체 구조",
    "제1장 · §§1.1, 1.4, 1.5",
    "측정값으로 현재를 갱신하고, 미래 입력 중 첫 값만 실제 공정에 줍니다",
)
for i, (title, body) in enumerate(
    [
        ("상태·외란 추정", "측정 y와 직전 입력으로\n현재 상태를 추정"),
        ("정상상태 목표", "기준값 r과 외란 추정으로\n목표 상태·입력을 계산"),
        ("예측·최적화", "추정 상태에서 미래를 예측\n첫 최적 입력만 선택"),
        ("실제 공정", "입력 u를 적용한 뒤\n다음 측정 y를 획득"),
    ]
):
    x = 0.75 + i * 3.13
    rect(s, x, 2.9, 2.65, 2.0, WHITE)
    text(s, x + 0.15, 3.12, 2.35, 0.5, title, 19, TEAL, True)
    text(s, x + 0.15, 3.85, 2.35, 0.9, body, 15)
    if i < 3:
        text(s, x + 2.7, 3.55, 0.4, 0.5, "→", 22, ORANGE, True)
rect(s, 1.05, 5.5, 11.15, 0.035, TEAL)
text(s, 0.8, 5.25, 0.45, 0.5, "↑", 24, TEAL, True)
text(s, 12.05, 5.13, 0.45, 0.5, "↓", 24, TEAL, True)
text(
    s,
    3.25,
    5.75,
    7.3,
    0.6,
    "새 측정과 적용한 입력을 기록하고 다음 시점에서 반복",
    18,
    TEAL,
    True,
)
notes(
    s,
    "상태 추정은 과거 측정의 요약, 목표 계산은 실행 가능한 정상 운전점 선택, 제어는 미래 입력의 최적화입니다. 추정 상태는 제어 최적화에도 전달됩니다.",
)
for d in DATA:
    s = slide(f"{d['n']}  {d['ko']}", "교재 해설", f"제1장 §{d['n']} · 인쇄 pp.{d['p']}")
    columns(s, d["points"], d["eqs"])
    notes(s, d["body"])
    for item in d["extras"]:
        if item["kind"] == "figure":
            figure(d, item)
        else:
            s = derivation(
                item["title"],
                f"제1장 §{d['n']} · 인쇄 pp.{d['p']}",
                item["subtitle"],
                item["rows"],
            )
            notes(s, d["body"])
s = slide(
    "노트북에서 가정을 바꿔 확인하세요",
    "실행과 참고",
    "제1장 · 인쇄 pp.1–88",
    "Rawlings · Mayne · Diehl, 2nd edition, 1st printing (2017)",
)
columns(
    s,
    [
        "환경 준비 후 원하는 실습 셀을 실행할 수 있습니다.",
        "그래프를 읽고 수치 검사의 의미까지 확인하세요.",
        "기본 결과는 깨끗한 로컬 커널에서 실행 검증했습니다.",
    ],
    [
        ("실행 위치", "GitHub README의 Colab 버튼"),
        ("수정 원본", "PPT 본문·수식 편집 가능"),
        ("재생성", "노트북 → 그림 → PPT → PDF"),
    ],
)
link = text(s, 0.85, 6.5, 11.5, 0.38, "Colab에서 제1장 노트북 열기", 17, TEAL, True)
link.text_frame.paragraphs[0].runs[0].hyperlink.address = COLAB
link.click_action.hyperlink.address = COLAB
prs.save(CH / "chapter01_slides.pptx")
print(f"Created {len(prs.slides)} slides: {CH}")
# Review metadata stays beside intermediate renders, not in chapter deliverables.
if args.manifest:
    args.manifest.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
